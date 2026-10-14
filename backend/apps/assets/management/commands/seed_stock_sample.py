"""Load an explicitly selected stock export into the isolated demo database."""

from datetime import date, datetime
from pathlib import Path
from uuid import UUID

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from rest_framework.exceptions import ValidationError

from apps.assets.management.commands.seed_demo import assert_demo_database, stable_id
from apps.assets.models import Asset, AssetType
from apps.assets.validation import FieldDefinition, JSONValue, default_stock_schema, validate_custom_data
from apps.tenants.models import Tenant
from apps.user.models import User


STOCK_TENANT_SLUG = "demo-stock-sample"
STOCK_TENANT_NAME = "Örnek Stok Tersanesi"
STOCK_TYPE_NAME = "İskele malzemeleri"
STOCK_CATALOG_CODE = "DEMO-STOCK-A-01"
STOCK_TYPE_ID = stable_id(f"type:{STOCK_CATALOG_CODE}")
SOURCE_HEADERS = (
    "ID", "Tersane", "Oluşturulma Tarihi", "Seri Numarası", "Adet", "İsim",
)
SAMPLE_DESCRIPTION = "Excel stok örneği; yalnız izole test ortamı için oluşturulmuştur."


def sample_stock_schema() -> list[FieldDefinition]:
    return default_stock_schema() + [
        {"key": "source_record_id", "label": "Kaynak kayıt kimliği", "type": "text"},
        {"key": "source_created_at", "label": "Kaynak kayıt tarihi", "type": "text"},
        {"key": "original_shipyard", "label": "Kaynak tersane", "type": "text"},
    ]


def stock_record_id(material_code: str, tenant_slug: str = STOCK_TENANT_SLUG) -> UUID:
    return stable_id(f"asset:stock-sample:{tenant_slug}:{STOCK_TYPE_ID}:{material_code}")


def source_text(value: object) -> str:
    """Keep exported text intact; Excel dates retain their original wall time."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def read_stock_export(path: Path) -> list[dict[str, JSONValue]]:
    """Validate every stock line before any database mutation is possible."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise CommandError("Excel dosyası açılamadı; mevcut veriler değiştirilmedi.") from exc
    try:
        if len(workbook.worksheets) != 1:
            raise CommandError("Stok örneği tek çalışma sayfası içermelidir.")
        sheet = workbook.worksheets[0]
        if sheet.max_column != len(SOURCE_HEADERS):
            raise CommandError("Beklenen altı Excel sütunu bulunamadı.")
        rows = sheet.iter_rows()
        header = next(rows, ())
        if tuple(cell.value for cell in header) != SOURCE_HEADERS:
            raise CommandError("Başlıklar ID, Tersane, Oluşturulma Tarihi, Seri Numarası, Adet ve İsim olmalıdır.")
        source_yard = None
        material_codes = set()
        records = []
        for row_number, cells in enumerate(rows, start=2):
            values = [cell.value for cell in cells]
            if all(value is None or value == "" for value in values):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise CommandError(f"{row_number}. satır formül içeriyor; stok aktarımı için hücre değerlerini kullanın.")
            source_id, yard, created_at, code, quantity, name = values
            if not isinstance(yard, str) or not yard.strip():
                raise CommandError(f"{row_number}. satırda kaynak tersane eksik.")
            normalized_yard = yard.strip()
            if source_yard is None:
                source_yard = normalized_yard
            elif normalized_yard != source_yard:
                raise CommandError("Dosya birden fazla tersane içeriyor; tersaneleri ayrı dosyalarda aktarın.")
            data = {
                "material_code": code,
                "material_name": name,
                "quantity": quantity,
                "source_record_id": source_text(source_id),
                "source_created_at": source_text(created_at),
                "original_shipyard": yard,
            }
            try:
                data = validate_custom_data(sample_stock_schema(), data, "stock")
            except ValidationError as exc:
                raise CommandError(f"{row_number}. satır geçersiz: {exc.detail}") from exc
            if data["material_code"] in material_codes:
                raise CommandError(f"{row_number}. satırda malzeme kodu tekrarlanıyor: {data['material_code']}")
            material_codes.add(data["material_code"])
            records.append(data)
        if not records:
            raise CommandError("Dosyada aktarılacak stok satırı bulunamadı.")
        return records
    finally:
        workbook.close()


class Command(BaseCommand):
    help = "Seçilen stok Excel'ini ayrı, paylaşıma kapalı bir demo tersanesine ekler."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Kaynak stok Excel dosyasının yolu.")
        parser.add_argument("--tenant-name", default=None)
        parser.add_argument("--tenant-slug", default=STOCK_TENANT_SLUG, help="Önceden hazırlanmış örneği güncellerken mevcut tersane kodunu kullanın.")
        parser.add_argument("--type-name", default=STOCK_TYPE_NAME)

    def handle(self, *args, **options):
        assert_demo_database()
        records = read_stock_export(Path(options["file"]))
        tenant_slug = options["tenant_slug"].strip()
        requested_name = options["tenant_name"]
        existing_name = Tenant.objects.filter(slug=tenant_slug).values_list("name", flat=True).first()
        tenant_name = requested_name.strip() if requested_name is not None else STOCK_TENANT_NAME
        if requested_name is None:
            if existing_name is not None:
                tenant_name = existing_name
        type_name = options["type_name"].strip()
        if not tenant_name or not type_name:
            raise CommandError("Örnek tersane ve tür adları boş bırakılamaz.")
        should_mark_example = existing_name is None or (requested_name is not None and tenant_name != existing_name)
        if should_mark_example and "örnek" not in tenant_name.lower():
            tenant_name += " — örnek stok"
        if len(tenant_name) > 255 or len(type_name) > 255:
            raise CommandError("Örnek tersane ve tür adları 1–255 karakter olmalıdır.")
        try:
            Tenant._meta.get_field("slug").clean(tenant_slug, None)
        except DjangoValidationError as exc:
            raise CommandError("Örnek tersane kodu geçersiz.") from exc
        if not tenant_slug:
            raise CommandError("Örnek tersane kodu boş bırakılamaz.")
        try:
            with transaction.atomic():
                tenant = self.get_tenant(tenant_name, tenant_slug)
                asset_type = self.get_type(type_name)
                admin = User.objects.filter(username="demo_admin", role="platform_admin").first()
                created_count = 0
                for data in records:
                    record, created = Asset.objects.get_or_create(
                        id=stock_record_id(data["material_code"], tenant_slug),
                        defaults={
                            "tenant": tenant,
                            "asset_type": asset_type,
                            "custom_data": data,
                            "created_by": admin,
                            "is_shared": False,
                        },
                    )
                    if record.tenant_id != tenant.pk or record.asset_type_id != asset_type.pk:
                        raise CommandError("Stok örneği kimliği başka bir kayda ait; mevcut veriler değiştirilmedi.")
                    created_count += int(created)
        except IntegrityError as exc:
            raise CommandError("Örnek tersane veya tür kimliği mevcut kayıtla çakışıyor; hiçbir kayıt değiştirilmedi.") from exc
        quantity = sum(record["quantity"] for record in records)
        self.stdout.write(self.style.SUCCESS(
            f"Örnek stok hazır: {tenant_name}; dosyada {len(records)} malzeme kalemi, {quantity} adet. "
            f"{created_count} yeni özel stok kaydı eklendi. Mevcut kayıtlar ve paylaşım kararları korundu."
        ))

    def get_tenant(self, name: str, slug: str = STOCK_TENANT_SLUG) -> Tenant:
        tenant, created = Tenant.objects.get_or_create(
            slug=slug,
            defaults={"name": name, "description": SAMPLE_DESCRIPTION},
        )
        if not created and (tenant.name != name or tenant.description != SAMPLE_DESCRIPTION):
            raise CommandError("Örnek tersane kimliği eşleşmiyor; mevcut veriler değiştirilmedi.")
        return tenant

    def get_type(self, name: str) -> AssetType:
        admin = User.objects.filter(username="demo_admin", role="platform_admin").first()
        schema = sample_stock_schema()
        asset_type, created = AssetType.objects.get_or_create(
            id=STOCK_TYPE_ID,
            defaults={
                "name": name,
                "catalog_code": STOCK_CATALOG_CODE,
                "category_code": "A",
                "sharing_class": "B",
                "tracking_mode": "stock",
                "schema": schema,
                "description": SAMPLE_DESCRIPTION,
                "created_by": admin,
            },
        )
        if not created and (
            asset_type.name != name
            or asset_type.catalog_code != STOCK_CATALOG_CODE
            or asset_type.tracking_mode != "stock"
            or asset_type.schema != schema
        ):
            raise CommandError("Örnek stok türünün kimliği veya alanları eşleşmiyor; mevcut veriler değiştirilmedi.")
        return asset_type
