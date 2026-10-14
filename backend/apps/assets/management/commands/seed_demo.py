"""Create fictional practice records only inside the isolated demo database."""

import json
import os
import uuid
from copy import deepcopy
from pathlib import Path

import openpyxl
from allauth.account.models import EmailAddress
from django.conf import settings
from django.contrib.auth.password_validation import validate_password
from django.contrib.sites.models import Site
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import connection, transaction
from openpyxl.styles import Alignment, Font, PatternFill

from apps.assets.catalog_import import HEADERS
from apps.assets.models import Asset, AssetType
from apps.assets.serializers import AssetSharingSerializer
from apps.notifications.models import Notification
from apps.tenants.models import Tenant
from apps.user.models import User, UserRole

DEMO_PASSWORD = "GisasDemo!2026"
DEMO_NAMESPACE = uuid.UUID("6d7b88e0-cf96-5b1e-bd75-e04b02f9237e")
DEMO_SCHEMA = [
    {"key": "asset_code", "label": "Ekipman kodu", "type": "text", "required": True, "is_unique_key": True},
    {"key": "brand", "label": "Marka", "type": "text"},
    {"key": "model", "label": "Model", "type": "text"},
    {"key": "serial_number", "label": "Seri numarası", "type": "text"},
    {"key": "capacity", "label": "Kapasite ve birimi", "type": "text"},
    {"key": "inspection_date", "label": "Son kontrol tarihi", "type": "date"},
    {"key": "in_service", "label": "Kullanıma hazır", "type": "checkbox"},
    {"key": "notes", "label": "Notlar", "type": "text"},
]
DEMO_TYPES = [
    ("DEMO-A-01", "Modüler iskele", "A", "A", "", "120 m²"),
    ("DEMO-B-01", "Mobil vinç", "B", "B", "", "25 ton"),
    ("DEMO-C-01", "MIG/MAG kaynak makinesi", "C", "A", "", "350 A"),
    ("DEMO-D-01", "Plazma kesim makinesi", "D", "B", "", "40 mm"),
    ("DEMO-E-01", "Mobil jeneratör", "E", "A", "", "100 kVA"),
    ("DEMO-F-01", "Vidalı kompresör", "F", "A", "", "10 bar"),
    ("DEMO-G-01", "Forklift", "G", "B", "", "3 ton"),
    ("DEMO-H-01", "Kumlama ünitesi", "H", "B", "", "200 litre"),
    ("DEMO-I-01", "Ultrasonik kalınlık ölçer", "I", "A", "", "0,8–300 mm"),
    ("DEMO-I-02", "Radyografik test cihazı", "I", "C", "Radyografi ekipmanı: kapsam dışı demo örneği.", "Demo RT"),
    ("DEMO-J-01", "Hidrolik pres", "J", "B", "", "50 ton"),
    ("DEMO-K-01", "Akülü matkap", "K", "A", "", "18 V"),
    ("DEMO-L-01", "Taşınabilir gaz ölçüm cihazı", "L", "A", "", "4 gaz"),
    ("DEMO-L-02", "Kişisel koruyucu donanım seti", "L", "C", "Kişisel koruyucu donanım: kişiye özel kullanım.", "1 set"),
    ("DEMO-M-01", "Yağ emici bariyer", "M", "B", "", "25 metre"),
]
DEMO_ACCOUNTS = [
    ("demo_admin", "admin@demo.gisas.test", "Deniz", "Platform", UserRole.PLATFORM_ADMIN, None, True),
    ("demo_yonetici", "yonetici@demo.gisas.test", "Ece", "Yönetici", UserRole.TENANT_ADMIN, 0, False),
    ("demo_personel", "personel@demo.gisas.test", "Mert", "Personel", UserRole.TENANT_USER, 0, False),
    ("demo_diger", "diger@demo.gisas.test", "Selin", "Poyraz", UserRole.TENANT_ADMIN, 1, False),
]


def stable_id(value):
    return uuid.uuid5(DEMO_NAMESPACE, value)


def assert_demo_database():
    """Reject production settings and a demo flag pointing at another database."""
    if not getattr(settings, "DEMO_MODE", False):
        raise CommandError("Bu komut yalnız config.demo_settings ile çalışır; mevcut veriler değiştirilmedi.")
    if connection.vendor != "sqlite":
        raise CommandError("Demo verisi yalnız izole SQLite veritabanında oluşturulabilir.")
    name = str(connection.settings_dict["NAME"])
    if connection.creation.is_in_memory_db(name) and getattr(settings, "DEMO_ALLOW_IN_MEMORY_SEED", False):
        return
    expected_path = getattr(settings, "DEMO_DATABASE_PATH", None)
    if not expected_path or Path(name).resolve() != Path(expected_path).resolve():
        raise CommandError("Veritabanı demo yoluyla eşleşmiyor; mevcut veriler değiştirilmedi.")


def write_workbook(path, rows, title="Katalog", guide=None, asset_type_id=None):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2" if title == "Katalog" else "A4"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = 28
    if guide:
        instructions = workbook.create_sheet("Rehber")
        for text in guide:
            instructions.append([text])
        instructions.column_dimensions["A"].width = 110
    if asset_type_id:
        mapping = workbook.create_sheet("_mapping")
        mapping.sheet_state = "hidden"
        mapping.append([json.dumps({field["label"]: field["key"] for field in DEMO_SCHEMA}, ensure_ascii=False)])
        mapping.append([str(asset_type_id)])
    workbook.save(path)
    workbook.close()


def write_samples(folder):
    folder.mkdir(parents=True, exist_ok=True)
    headers = list(HEADERS)
    guide = [
        "Bu dosya yalnız kurmaca demo verisi içerir.",
        "Varlık Tipleri > Katalogdan aktar: dosyayı yükleyin, önizleyin ve onaylayın.",
        "Her satır bir varlık tipidir; fiziksel ekipmanlar Envanter ekranından eklenir.",
        "Ön paylaşım sınıfı saha onayı veya paylaşım izni değildir.",
    ]
    write_workbook(folder / "katalog-gecerli.xlsx", [headers,
        ["DEMO-PRACTICE-B-01", "Pratik — Portatif vinç", "B", "B", "", "Katalog aktarımı alıştırması: 2 ton örnek kapasite."],
        ["DEMO-PRACTICE-E-01", "Pratik — Sessiz jeneratör", "E", "A", "", "Katalog aktarımı alıştırması: 40 kVA örnek kapasite."],
        ["DEMO-PRACTICE-K-01", "Pratik — El taşlama makinesi", "K", "A", "", "K kategorisi otomatik kapsam dışıdır."],
    ], guide=guide)
    write_workbook(folder / "katalog-hatali.xlsx", [headers,
        ["", "Hatalı — Kodsuz tip", "B", "B", "", "Katalog kodu eksik."],
        ["DEMO-ERROR-01", "Hatalı — Geçersiz kategori", "Z", "B", "", "Kategori A–M aralığında olmalı."],
        ["DEMO-ERROR-01", "Hatalı — Tekrarlanan kod", "B", "B", "", "Aynı katalog kodu iki satırda kullanılıyor."],
    ], guide=[*guide, "Bu dosya bilerek hatalıdır. Önizlemede satır hataları görülür; hiçbir kayıt aktarılmaz."])
    write_workbook(folder / "katalog-guncelleme.xlsx", [headers,
        ["DEMO-B-01", "Mobil vinç", "B", "B", "", "Demo güncelleme alıştırması: operatör ve taşıma koşulları saha değerlendirmesinde incelenir."],
    ], guide=[*guide, "Katalog koduna göre güncelle modunu seçin. Mobil vinç açıklaması güncellenir; ekipmanlar korunur."])
    field_headers = [field["label"] for field in DEMO_SCHEMA]
    info = ["Zorunlu / benzersiz", "Metin", "Metin", "Metin", "Kapasite ve birimi", "YYYY-AA-GG", "Evet/Hayır", "Metin"]
    example = ["ÖRNEK-SATIR-AKTARILMAZ", "DemoMarka", "FL-30", "ÖRNEK", "3 ton", "2026-09-15", True, "Örnek satır aktarılmaz."]
    rows = [field_headers, info, example,
        ["PRACTICE-FL-001", "DemoMarka", "FL-30", "DM-IMP-001", "3 ton", "2026-09-15", True, "Envanter aktarımı alıştırması."],
        ["PRACTICE-FL-002", "Örnek Endüstri", "FL-25", "DM-IMP-002", "2,5 ton", "2026-08-20", False, "Kontrol bekleyen demo ekipmanı."],
    ]
    inventory_guide = [
        "Bu dosya yalnız kurmaca demo verisi içerir.",
        "Envanter > Excel'den aktar > Forklift tipini seçin > Şablonlu aktarım.",
        "İlk üç satır başlık, alan bilgisi ve örnektir. Ekipmanlar dördüncü satırdan aktarılır.",
        "Dosya Forklift tipine aittir. Ekipman kodları kurmaca alıştırma kayıtlarını tanımlar.",
    ]
    forklift_id = stable_id("type:DEMO-G-01")
    write_workbook(folder / "envanter-forklift.xlsx", rows, title="Forklift", guide=inventory_guide, asset_type_id=forklift_id)
    write_workbook(folder / "envanter-forklift-hatali.xlsx", [field_headers, info, example,
        ["", "DemoMarka", "FL-30", "DM-ERR-001", "3 ton", "2026-09-15", True, "Zorunlu ekipman kodu eksik."],
        ["PRACTICE-FL-ERROR", "DemoMarka", "FL-30", "DM-ERR-002", "3 ton", "dün", "belki", "Tarih ve evet/hayır hücreleri hatalı."],
    ], title="Forklift", guide=[*inventory_guide, "Bilerek hatalı satırlar: ekipman kodu, tarih ve evet/hayır alanlarını düzeltin."], asset_type_id=forklift_id)


class Command(BaseCommand):
    help = "İzole demo veritabanına kurmaca tersane, kullanıcı, katalog ve ekipman örnekleri ekler."

    def add_arguments(self, parser):
        parser.add_argument("--write-samples", action="store_true", help="Demo Excel alıştırmalarını oluştur.")
        parser.add_argument(
            "--model-count-demo", action="store_true",
            help="Mavi Körfez'e aynı modelden 50 özel fiziksel ekipman örneği ekle.",
        )

    def handle(self, *args, **options):
        assert_demo_database()
        password = os.environ.get("DEMO_PASSWORD", DEMO_PASSWORD)
        try:
            validate_password(password)
        except ValidationError as exc:
            raise CommandError("Demo parolası geçersiz: " + "; ".join(exc.messages)) from exc
        with transaction.atomic():
            yards = self.seed_yards()
            users = self.seed_users(yards, password)
            types = self.seed_types(users[0])
            self.seed_assets(yards, types, users)
            self.seed_shared_assets(yards, types, users[0])
            if options["model_count_demo"]:
                self.seed_model_count_assets(types, users[1])
            self.seed_notifications(users)
            Site.objects.get_or_create(pk=settings.SITE_ID, defaults={"domain": "localhost", "name": "GİSAŞ Demo"})
        if options["write_samples"]:
            write_samples(Path(settings.DEMO_FILES_DIR))
        equipment_count = 83 if options["model_count_demo"] else 33
        self.stdout.write(self.style.SUCCESS(
            f"Demo hazır: 3 kurmaca tersane, 4 hesap, 15 varlık tipi, {equipment_count} başlangıç ekipmanı, 12 bildirim. "
            "Tekrar çalıştırma mevcut kayıtları ve parolaları değiştirmez."
        ))

    def seed_yards(self):
        yards = []
        for slug, name, region in [
            ("demo-mavi-korfez", "Mavi Körfez Demo Tersanesi", "Tuzla"),
            ("demo-poyraz", "Poyraz Demo Tersanesi", "Yalova"),
            ("demo-cinar", "Çınar Demo Tersanesi", "Kocaeli"),
        ]:
            yard, _ = Tenant.objects.get_or_create(slug=slug, defaults={
                "name": name,
                "description": "Eğitim ve arayüz testleri için oluşturulmuş kurmaca tersane.",
                "address": f"Demo Eğitim Alanı — {region}",
                "email": f"{slug}@gisas.test",
            })
            yards.append(yard)
        return yards

    def seed_users(self, yards, password):
        users = []
        for username, email, first_name, last_name, role, yard_index, platform in DEMO_ACCOUNTS:
            tenant = None if yard_index is None else yards[yard_index]
            user, created = User.objects.get_or_create(username=username, defaults={
                "email": email, "first_name": first_name, "last_name": last_name,
                "role": role, "tenant": tenant, "is_staff": platform, "is_superuser": platform,
                "preferred_language": "tr",
            })
            if not created and (user.email != email or user.role != role or user.tenant_id != getattr(tenant, "pk", None)):
                raise CommandError(f"{username} demo kimliğiyle eşleşmiyor; hiçbir kayıt değiştirilmedi.")
            if created:
                user.set_password(password)
                user.save(update_fields=["password"])
            EmailAddress.objects.get_or_create(user=user, email=email, defaults={"verified": True, "primary": True})
            users.append(user)
        return users

    def seed_types(self, admin):
        types = []
        for code, name, category, grade, reason, capacity in DEMO_TYPES:
            asset_type, _ = AssetType.objects.get_or_create(catalog_code=code, defaults={
                "id": stable_id(f"type:{code}"), "name": name,
                "category_code": category, "sharing_class": grade, "exclusion_reason": reason,
                "description": f"Kurmaca demo katalog kaydı. Örnek kapasite: {capacity}. Saha değerlendirmesi gerekir.",
                "schema": deepcopy(DEMO_SCHEMA), "created_by": admin,
            })
            types.append(asset_type)
        return types

    def seed_assets(self, yards, types, users):
        for index, asset_type in enumerate(types):
            capacity = DEMO_TYPES[index][-1]
            for copy_index in range(2):
                yard_index = (index + copy_index) % len(yards)
                yard = yards[yard_index]
                code = f"{yard.slug.removeprefix('demo-').upper()}-{asset_type.category_code}-{index + 1:02d}-{copy_index + 1}"
                assignee = users[2] if yard_index == 0 else users[3] if yard_index == 1 else None
                Asset.objects.get_or_create(id=stable_id(f"asset:{index}:{copy_index}"), defaults={
                    "tenant": yard, "asset_type": asset_type, "created_by": users[0], "assigned_to": assignee,
                    "custom_data": {
                        "asset_code": code, "brand": "DemoMarka" if copy_index == 0 else "Örnek Endüstri",
                        "model": f"DM-{asset_type.category_code}-{index + 1:02d}",
                        "serial_number": f"DEMO-SN-{index + 1:02d}-{copy_index + 1}",
                        "capacity": capacity, "inspection_date": "2026-09-15" if copy_index == 0 else "2026-08-20",
                        "in_service": index % 5 != 0 or copy_index == 0,
                        "notes": "Kurmaca test kaydı. Paylaşım için saha onayı gereklidir.",
                    },
                })

    def seed_model_count_assets(self, types, manager):
        forklift = next(asset_type for asset_type in types if asset_type.catalog_code == "DEMO-G-01")
        for number in range(1, 51):
            Asset.objects.get_or_create(id=stable_id(f"asset:model-count-50:{number}"), defaults={
                "tenant": manager.tenant, "asset_type": forklift, "created_by": manager,
                "is_shared": False,
                "custom_data": {
                    "asset_code": f"DEMO-FL-50-{number:03d}", "brand": "Toyota", "model": "8FD",
                    "serial_number": f"DEMO-SERIAL-{number:03d}", "capacity": "3 ton",
                    "inspection_date": "2026-09-15", "in_service": True,
                    "notes": "Kurmaca test kaydı. Aynı modelden 50 fiziksel ekipmanı inceleme örneği.",
                },
            })

    def seed_notifications(self, users):
        notices = [
            ("welcome", "Demo ortamına hoş geldiniz", "Kullanım rehberinden katalog, envanter ve kullanıcı akışlarını deneyebilirsiniz.", "info"),
            ("catalog", "Katalog alıştırmaları hazır", "Geçerli ve hatalı örnek Excel dosyaları Kullanım rehberi ekranında bulunuyor.", "success"),
            ("review", "Saha değerlendirmesi örneği", "A/B/C ön paylaşım sınıflarını inceleyin; gerçek ekipman onayı için saha değerlendirmesi gerekir.", "warning"),
        ]
        for user in users:
            for key, title, message, kind in notices:
                marker = f"demo:{user.username}:{key}"
                Notification.objects.get_or_create(user=user, content__demo_key=marker, defaults={
                    "content": {"demo_key": marker, "title": title, "message": message, "type": kind},
                })

    def seed_shared_assets(self, yards, types, admin):
        examples = [
            (0, 'DEMO-G-01', '3 ton forklift. Taşıma ve kullanım koşulları tersane ile görüşülür.'),
            (1, 'DEMO-B-01', '25 ton mobil vinç. Operatör ve taşıma koşulları ayrıca değerlendirilir.'),
            (2, 'DEMO-J-01', '50 ton hidrolik pres. Yerinde kullanım koşulları tersane ile görüşülür.'),
        ]
        by_code = {asset_type.catalog_code: asset_type for asset_type in types}
        for yard_index, catalog_code, public_note in examples:
            asset_type = by_code[catalog_code]
            asset, created = Asset.objects.get_or_create(id=stable_id(f'shared-example:{catalog_code}'), defaults={
                'tenant': yards[yard_index], 'asset_type': asset_type, 'created_by': admin,
                'custom_data': {
                    'asset_code': f'DEMO-SHARED-{yard_index + 1:02d}', 'brand': 'DemoMarka',
                    'notes': 'Kurum içi demo notu. Ortak listede gösterilmez.',
                },
            })
            if created and asset.tenant.is_active and not asset_type.sharing_excluded:
                serializer = AssetSharingSerializer(asset, data={'is_shared': True, 'sharing_note': public_note})
                serializer.is_valid(raise_exception=True)
                serializer.save()
