from io import BytesIO
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from rest_framework.test import APITestCase

from apps.tenants.models import Tenant
from apps.user.models import User
from .catalog_import import HEADERS
from .models import Asset, AssetType


class CatalogImportTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create(username='catalog-admin', role='platform_admin')
        self.client.force_authenticate(self.admin)
        self.rows = [['B-001', 'Portal vinç', 'B', 'B', '', 'Örnek açıklama']]
        self.content = self.excel(self.rows)

    @staticmethod
    def excel(rows, headers=None):
        workbook = openpyxl.Workbook()
        workbook.active.title = 'Katalog'
        workbook.active.append(headers or list(HEADERS))
        for row in rows:
            workbook.active.append(row)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def post(self, action, content=None, mode='create', **kwargs):
        return self.client.post(f'/api/assets/types/catalog_{action}/', {
            'file': SimpleUploadedFile('catalog.xlsx', content or self.content),
            'mode': mode, **kwargs,
        }, format='multipart')

    def test_preview_is_read_only_and_confirm_creates_usable_type(self):
        preview = self.post('preview')
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(AssetType.objects.count(), 0)
        self.assertEqual(preview.data['summary']['create'], 1)
        result = self.post('import', preview_token=preview.data['preview_token'])
        self.assertEqual(result.status_code, 200, result.data)
        created = AssetType.objects.get(catalog_code='B-001')
        self.assertEqual(created.created_by, self.admin)
        self.assertTrue(created.schema[0]['is_unique_key'])
        self.assertFalse(created.sharing_excluded)
        # A fresh preview of the same workbook never duplicates the type.
        repeated = self.post('preview')
        self.assertEqual(repeated.data['summary']['skip'], 1)
        self.assertEqual(self.post('import', preview_token=repeated.data['preview_token']).status_code, 200)
        self.assertEqual(AssetType.objects.count(), 1)

    def test_update_preserves_schema_creator_and_inventory(self):
        creator = User.objects.create(username='original')
        item = AssetType.objects.create(name='Old crane', catalog_code='B-001', schema=[
            {'key': 'serial', 'label': 'Serial', 'type': 'text'},
        ], created_by=creator)
        tenant = Tenant.objects.create(name='Yard')
        asset = Asset.objects.create(tenant=tenant, asset_type=item, custom_data={'serial': 'ABC'})
        preview = self.post('preview', mode='update')
        self.assertEqual(preview.data['summary']['update'], 1)
        result = self.post('import', mode='update', preview_token=preview.data['preview_token'])
        self.assertEqual(result.status_code, 200, result.data)
        item.refresh_from_db()
        asset.refresh_from_db()
        self.assertEqual(item.name, 'Portal vinç')
        self.assertEqual(item.created_by, creator)
        self.assertEqual(item.schema[0]['key'], 'serial')
        self.assertEqual(asset.custom_data, {'serial': 'ABC'})

    def test_errors_block_entire_file(self):
        content = self.excel(self.rows + [['B-001', 'Duplicate', 'Z', 'D', '', '']])
        preview = self.post('preview', content)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data['summary']['error'], 1)
        self.assertIsNone(preview.data['preview_token'])
        self.assertGreater(len(preview.data['rows'][1]['errors']), 1)
        self.assertEqual(self.post('import', content).status_code, 400)
        self.assertFalse(AssetType.objects.exists())

    def test_transaction_rolls_back_on_conflict(self):
        preview = self.post('preview')
        from .serializers import AssetTypeSerializer
        original = AssetTypeSerializer.create
        def fail_after_create(serializer, data):
            original(serializer, data)
            raise IntegrityError('simulated race')
        with patch.object(AssetTypeSerializer, 'create', fail_after_create):
            result = self.post('import', preview_token=preview.data['preview_token'])
        self.assertEqual(result.status_code, 409)
        self.assertFalse(AssetType.objects.exists())

    def test_stale_or_changed_preview_cannot_overwrite(self):
        preview = self.post('preview')
        token = preview.data['preview_token']
        self.assertEqual(self.post('import', mode='update', preview_token=token).status_code, 400)
        changed = self.excel([['B-002', 'Different', 'B', 'B', '', '']])
        self.assertEqual(self.post('import', changed, preview_token=token).status_code, 400)
        AssetType.objects.create(name='Concurrent edit')
        self.assertEqual(self.post('import', preview_token=token).status_code, 409)
        self.assertEqual(AssetType.objects.count(), 1)

    def test_member_cannot_import_or_download_catalog(self):
        self.client.force_authenticate(User.objects.create(username='member'))
        self.assertEqual(self.post('preview').status_code, 403)
        self.assertEqual(self.post('import').status_code, 403)
        self.assertEqual(self.client.get('/api/assets/types/catalog_template/').status_code, 403)

    def test_template_has_guidance_but_no_importable_examples(self):
        response = self.client.get('/api/assets/types/catalog_template/')
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(list(next(workbook['Katalog'].values)), list(HEADERS))
        self.assertIn('Rehber', workbook.sheetnames)
        self.assertEqual(workbook['Katalog'].max_row, 1)
        workbook.close()
        self.assertEqual(self.post('preview', response.content).status_code, 400)

    def test_malformed_workbooks_and_formulas_are_rejected(self):
        self.assertEqual(self.post('preview', b'not excel').status_code, 400)
        self.assertEqual(self.post('preview', self.excel(self.rows, ['Wrong'])).status_code, 400)
        formula = self.excel([['B-001', '=1+1', 'B', 'B', '', '']])
        self.assertEqual(self.post('preview', formula).data['summary']['error'], 1)

    def test_category_class_and_explicit_exclusion_are_independent(self):
        for category, grade, reason, excluded in [('K', 'A', '', True), ('I', 'B', 'Radyografi', True),
                                                  ('L', 'A', 'KKD', True), ('B', 'C', '', True), ('B', 'A', '', False)]:
            with self.subTest(category=category, grade=grade):
                response = self.client.post('/api/assets/types/', {
                    'name': f'{category}-{grade}-{reason}', 'category_code': category,
                    'sharing_class': grade, 'exclusion_reason': reason, 'sharing_excluded': False, 'schema': [],
                }, format='json')
                self.assertEqual(response.status_code, 201, response.data)
                self.assertEqual(response.data['sharing_excluded'], excluded)

    def test_catalog_code_normalization_and_empty_legacy_codes(self):
        item = AssetType.objects.create(name='Original', catalog_code='B-001')
        response = self.client.post('/api/assets/types/', {'name': 'Duplicate', 'catalog_code': 'b-001'}, format='json')
        self.assertEqual(response.status_code, 400)
        for name in ['Legacy one', 'Legacy two']:
            response = self.client.post('/api/assets/types/', {'name': name, 'catalog_code': ''}, format='json')
            self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(AssetType.objects.filter(catalog_code__isnull=True).count(), 2)
        self.assertEqual(self.client.delete(f'/api/assets/types/{item.pk}/').status_code, 204)

    def test_type_with_inventory_cannot_be_deleted(self):
        item = AssetType.objects.create(name='Protected')
        Asset.objects.create(tenant=Tenant.objects.create(name='Yard'), asset_type=item)
        self.assertEqual(self.client.delete(f'/api/assets/types/{item.pk}/').status_code, 409)
        self.assertEqual(Asset.objects.count(), 1)

    def test_skipped_rows_show_current_classification(self):
        AssetType.objects.create(name='Portal vinç', catalog_code='B-001', category_code='B', sharing_class='C')
        preview = self.post('preview')
        row = preview.data['rows'][0]
        self.assertEqual(row['operation'], 'skip')
        self.assertEqual(row['skip_reason'], 'exists')
        self.assertEqual(row['current_values']['sharing_class'], 'C')
        self.assertTrue(row['sharing_excluded'])
        self.assertEqual(row['changes'], {})

    def test_update_preview_exposes_cleared_exclusion(self):
        AssetType.objects.create(name='Portal vinç', catalog_code='B-001', category_code='B',
                                 sharing_class='B', exclusion_reason='Bakım kısıtı', description='Eski açıklama')
        row = self.post('preview', mode='update').data['rows'][0]
        self.assertEqual(row['changes']['exclusion_reason'], {'before': 'Bakım kısıtı', 'after': ''})
        self.assertFalse(row['sharing_excluded'])

    def test_extra_columns_outside_preview_cannot_be_silently_ignored(self):
        content = self.excel([self.rows[0] + [None, 'Hidden value in column H']])
        response = self.post('preview', content)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse(AssetType.objects.exists())
