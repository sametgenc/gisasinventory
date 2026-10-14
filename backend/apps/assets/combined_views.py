"""Authenticated endpoints for reviewing and importing one workbook."""

from io import BytesIO

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.tenants.permissions import IsActiveTenantUser
from .combined_import import (
    FIELD_HEADERS, build_plan, commit_workbook, parse_decisions, preview_identity,
    read_workbook, resolve_tenant, signed_preview, verify_preview, WorkbookConflict,
)


class WorkbookInputMixin:
    permission_classes = [IsAuthenticated, IsActiveTenantUser]

    def workbook_input(self, request):
        tenant = resolve_tenant(request.user, request.data.get('tenant'))
        decisions = parse_decisions(request.data.get('decisions'))
        content, sheets = read_workbook(request.FILES.get('file'))
        return sheets, decisions, tenant, preview_identity(content, decisions, tenant, request.user)


class WorkbookPreviewView(WorkbookInputMixin, APIView):
    def post(self, request):
        sheets, decisions, tenant, identity = self.workbook_input(request)
        plan = build_plan(sheets, decisions, tenant, request)
        result = plan.response()
        result['preview_token'] = signed_preview(identity, plan.fingerprint) if result['can_commit'] else None
        return Response(result)


class WorkbookCommitView(WorkbookInputMixin, APIView):
    def post(self, request):
        sheets, decisions, tenant, identity = self.workbook_input(request)
        preview = verify_preview(request.data.get('preview_token'), identity)
        try:
            summary = commit_workbook(sheets, decisions, tenant, request, preview)
        except WorkbookConflict as exc:
            return Response({'detail': str(exc)}, status=409)
        return Response({'summary': summary, 'detail': 'Excel kayıtları aktarıldı. Envanter kendiliğinden paylaşıma açılmaz.'})


class WorkbookTemplateView(WorkbookInputMixin, APIView):
    def get(self, request):
        resolve_tenant(request.user, request.query_params.get('tenant'))
        workbook = openpyxl.Workbook()
        types = workbook.active
        types.title = 'Ekipman türleri'
        types.append(['Katalog kodu', 'Ekipman türü', 'Kategori kodu', 'Ön paylaşım sınıfı',
                      'Kapsam dışı nedeni', 'Açıklama', 'Takip biçimi'])
        fields = workbook.create_sheet('Alanlar')
        fields.append(list(FIELD_HEADERS))
        inventory = workbook.create_sheet('Envanter')
        inventory.append(['Katalog kodu', 'Envanter kodu'])
        for sheet in (types, fields, inventory):
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='25465C')
                sheet.column_dimensions[cell.column_letter].width = 30
        for sheet, column, options in [
            (types, 'C', 'A,B,C,D,E,F,G,H,I,J,K,L,M'), (types, 'D', 'A,B,C'),
            (types, 'G', 'individual,stock'), (fields, 'D', 'text,number,date,select,checkbox,email,phone'),
            (fields, 'E', 'Evet,Hayır'), (fields, 'F', 'identifier,display_name,quantity'), (fields, 'H', 'Evet,Hayır'),
        ]:
            validation = DataValidation(type='list', formula1=f'"{options}"', allow_blank=True)
            validation.showErrorMessage = True
            validation.errorTitle = 'Geçersiz seçim'
            validation.error = 'Listeden bir değer seçin veya boş bırakın.'
            sheet.add_data_validation(validation)
            validation.add(f'{column}2:{column}2001')
        guide = workbook.create_sheet('Rehber')
        for row in [
            ['BİR EXCEL DOSYASIYLA AKTARIM'],
            ['1. Ekipman türleri', 'Ortak tür tanımıdır. Envanter veya adet değildir. Türü oluşturmak/düzenlemek platform yöneticisi yetkisi ister.'],
            ['2. Alanlar', 'Her satır bir türün alanıdır. Katalog koduyla bağlanır. Zorunlu boşsa yeni alan isteğe bağlıdır. Mevcut diğer alanlar korunur.'],
            ['3. Envanter', 'Her satır seçilen tersanenin fiziksel ekipmanı veya adetli malzeme kaydıdır. Hedef tersane uygulamada seçilir.'],
            ['Katalog kodu', 'Türün kalıcı ortak kimliğidir. Örn. ORNEK-001. A–M kategorileri katalog kodu değildir. Boş yeni tür kodu uygulamada önerilir.'],
            ['Envanter kodu', 'Tersanenin kendi kayıt kodudur. Aynı türde aynı kod güncellenir; başka tersanenin kaydı etkilenmez.'],
            ['Dinamik sütunlar', 'Envanter sayfasına Alan adı veya Alan kodu ile sütun ekleyin. Farklı türler için geçerli olmayan hücreler boş kalabilir.'],
            ['Alan kodu', 'Kalıcı alan adıdır. Örn. uretici. Boşsa alan adından önerilir.'],
            ['Görevi', 'identifier: benzersiz kayıt eşleştirme; display_name: görünen ad; quantity: adetli stok miktarı. Görev seçmek zorunlu yapmak değildir.'],
            ['Takip biçimi', 'individual: her ekipman ayrı satır; stock: aynı malzeme bir satır ve miktar alanı. Boş yeni tür tek tek başlar.'],
            ['Seçenekler', 'select alanlarında seçenekleri | ile ayırın. Örn. İyi|Bakımda|Kullanım dışı.'],
            ['Yalnızca envanter', 'Mevcut katalog kodlarıyla Envanter sayfasını doldurun. İlk iki sayfa boş kalabilir.'],
            ['Kontrol', 'Önce önizleyin. Tanınmayan sütunlar, eksik zorunlu alanlar veya belirsiz kararlar çözülmeden hiçbir kayıt aktarılmaz.'],
            ['Atlanan satırlar', 'Boş satırlar, tekrar başlıklar, birleştirilmiş bölüm başlıkları ve yalnızca sayısal toplam satırları aktarılmaz.'],
            ['Paylaşım', 'Aktarım paylaşım izni vermez. Ekipmanı sahibi ayrıca paylaşıma açar; kullanım talebini ekipman sahibi tersane yöneticisi onaylar.'],
            ['Bedel', 'Ortak kullanım ücretsizdir; ücret veya ödeme alanı yoktur.'],
            ['Kaynak taksonomi', '1-TAKSONOMI içeren GISBIR dosyası doğrudan tür önerisi verir. B/C satırlarında sınıfı seçmeniz gerekir.'],
        ]:
            guide.append(row)
        guide.column_dimensions['A'].width = 28
        guide.column_dimensions['B'].width = 115
        guide.freeze_panes = 'A2'
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="birlesik-envanter-sablonu.xlsx"'
        return response
