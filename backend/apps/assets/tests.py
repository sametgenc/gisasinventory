import json
from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils.http import parse_header_parameters
from rest_framework.test import APITestCase

from apps.tenants.models import Tenant
from apps.user.models import User
from .models import Asset, AssetType


def workbook_file(rows):
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return SimpleUploadedFile("assets.xlsx", buffer.getvalue())


class AssetAccessTests(APITestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(name="First yard")
        self.other = Tenant.objects.create(name="Other yard")
        self.user = User.objects.create(username="member", tenant=self.tenant)
        self.asset_type = AssetType.objects.create(name="Forklift", schema=[
            {"key": "serial", "label": "Serial", "type": "text", "required": True},
        ])
        self.asset = Asset.objects.create(tenant=self.tenant, asset_type=self.asset_type,
                                          custom_data={"serial": "F-1"})
        self.client.force_authenticate(self.user)

    def payload(self, **overrides):
        return {"asset_type": str(self.asset_type.pk), "custom_data": {"serial": "F-2"}, **overrides}

    def test_unassigned_member_cannot_write_into_explicit_tenant(self):
        self.user.tenant = None
        self.user.save()
        response = self.client.post("/api/assets/items/", self.payload(tenant=self.other.pk), format="json")
        self.assertIn(response.status_code, (400, 403))
        self.assertEqual(Asset.objects.count(), 1)

    def test_member_cannot_read_or_change_another_tenants_asset(self):
        self.asset.tenant = self.other
        self.asset.save()
        url = f"/api/assets/items/{self.asset.pk}/"
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.patch(url, {"custom_data": {"serial": "x"}}, format="json").status_code, 404)

    def test_disabled_tenant_cannot_access_assets(self):
        self.tenant.is_active = False
        self.tenant.save()
        self.assertEqual(self.client.get("/api/assets/items/").status_code, 403)

    def test_partial_update_preserves_required_data_and_tenant(self):
        self.user.is_superuser = True
        self.user.tenant = self.other
        self.user.save()
        response = self.client.patch(f"/api/assets/items/{self.asset.pk}/", {"assigned_to": None}, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.tenant_id, self.tenant.pk)
        self.assertEqual(self.asset.custom_data, {"serial": "F-1"})

    def test_assignee_must_belong_to_asset_tenant(self):
        other_user = User.objects.create(username="other", tenant=self.other)
        response = self.client.post("/api/assets/items/", self.payload(assigned_to=other_user.pk), format="json")
        self.assertEqual(response.status_code, 400, response.data)

    def test_invalid_dynamic_values_are_rejected(self):
        for field_type, value, extra in [
            ("date", "yesterday", {}), ("checkbox", "false", {}),
            ("select", "invalid", {"options": ["valid"]}),
            ("number", "NaN", {}), ("email", "not-email", {}),
        ]:
            with self.subTest(field_type=field_type):
                self.asset_type.schema = [{"key": "value", "label": "Value", "type": field_type, **extra}]
                self.asset_type.save()
                response = self.client.post("/api/assets/items/", self.payload(custom_data={"value": value}), format="json")
                self.assertEqual(response.status_code, 400, response.data)

    def test_custom_data_must_be_an_object(self):
        response = self.client.post("/api/assets/items/", self.payload(custom_data=[]), format="json")
        self.assertEqual(response.status_code, 400)

    def test_schema_rejects_duplicate_keys(self):
        self.user.is_superuser = True
        self.user.save()
        response = self.client.post("/api/assets/types/", {
            "name": "Duplicate", "schema": self.asset_type.schema * 2,
        }, format="json")
        self.assertEqual(response.status_code, 400)

    def test_raw_import_enforces_required_fields(self):
        response = self.client.post("/api/assets/items/raw_import/", {
            "asset_type": str(self.asset_type.pk), "mapping": json.dumps({"Serial": "serial"}),
            "file": workbook_file([["Serial", "Note"], [None, "not an empty row"]]),
        }, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["created_count"], 0)
        self.assertEqual(response.data["error_count"], 1)

    def test_raw_import_updates_by_unique_key_and_preserves_unmapped_data(self):
        self.asset_type.schema[0]["is_unique_key"] = True
        self.asset_type.schema.append({"key": "capacity", "label": "Capacity", "type": "number"})
        self.asset_type.save()
        self.asset.custom_data["capacity"] = 4
        self.asset.save()
        response = self.client.post("/api/assets/items/raw_import/", {
            "asset_type": str(self.asset_type.pk), "mapping": json.dumps({"Serial": "serial"}),
            "file": workbook_file([["Serial"], ["F-1"]]),
        }, format="multipart")
        self.assertEqual(response.data["updated_count"], 1, response.data)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.custom_data["capacity"], 4)

    def test_template_import_accepts_zero_and_false_and_records_creator(self):
        self.asset_type.schema = [
            {"key": "capacity", "label": "Capacity", "type": "number", "required": True},
            {"key": "active", "label": "Active", "type": "checkbox", "required": True},
        ]
        self.asset_type.save()
        response = self.client.post("/api/assets/items/bulk_import/", {
            "asset_type": str(self.asset_type.pk),
            "file": workbook_file([["Capacity", "Active"], ["info"], ["example"], [0, False]]),
        }, format="multipart")
        self.assertEqual(response.data["success_count"], 1, response.data)
        created = Asset.objects.exclude(pk=self.asset.pk).get()
        self.assertEqual(created.custom_data, {"capacity": 0, "active": False})
        self.assertEqual(created.created_by_id, self.user.pk)

    def test_invalid_import_mapping_returns_validation_error(self):
        response = self.client.post("/api/assets/items/raw_import/", {
            "asset_type": str(self.asset_type.pk), "mapping": '["serial"]',
            "file": workbook_file([["Serial"], ["F-2"]]),
        }, format="multipart")
        self.assertEqual(response.status_code, 400)

    def test_raw_import_rejects_invalid_checkbox_instead_of_silently_using_false(self):
        self.asset_type.schema = [{"key": "active", "label": "Active", "type": "checkbox"}]
        self.asset_type.save()
        response = self.client.post('/api/assets/items/raw_import/', {
            'asset_type': str(self.asset_type.pk), 'mapping': json.dumps({'Active': 'active'}),
            'file': workbook_file([['Active'], ['not a boolean']]),
        }, format='multipart')
        self.assertEqual(response.data['created_count'], 0, response.data)
        self.assertEqual(response.data['error_count'], 1)

    def test_template_import_keeps_blank_column_positions(self):
        response = self.client.post('/api/assets/items/bulk_import/', {
            'asset_type': str(self.asset_type.pk),
            'file': workbook_file([[None, 'Serial'], ['info'], ['example'], ['ignored', 'F-3']]),
        }, format='multipart')
        self.assertEqual(response.data['success_count'], 1, response.data)
        self.assertTrue(Asset.objects.filter(custom_data__serial='F-3').exists())

    def test_raw_import_rejects_mapping_for_a_nonexistent_header(self):
        response = self.client.post('/api/assets/items/raw_import/', {
            'asset_type': str(self.asset_type.pk), 'mapping': json.dumps({'Missing': 'serial'}),
            'file': workbook_file([['Serial'], ['F-3']]),
        }, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_repeated_template_import_updates_unique_key_without_duplicates(self):
        original_creator = User.objects.create(username="original_creator", tenant=self.tenant)
        self.asset_type.schema[0]["is_unique_key"] = True
        self.asset_type.schema.extend([
            {"key": "capacity", "label": "Capacity", "type": "number"},
            {"key": "active", "label": "Active", "type": "checkbox"},
            {"key": "notes", "label": "Notes", "type": "text"},
        ])
        self.asset_type.save()
        self.asset.custom_data.update({"capacity": 5, "active": True, "notes": "Keep this note"})
        self.asset.created_by = original_creator
        self.asset.assigned_to = self.user
        self.asset.save()
        foreign = Asset.objects.create(tenant=self.other, asset_type=self.asset_type,
                                       custom_data={"serial": "F-1", "capacity": 99})
        rows = [["Serial", "Capacity", "Active"], ["info"], ["example"], ["F-1", 0, False], ["F-2", 4, True]]
        for attempt in range(2):
            response = self.client.post("/api/assets/items/bulk_import/", {
                "asset_type": str(self.asset_type.pk), "file": workbook_file(rows),
            }, format="multipart")
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data["success_count"], 2)
            self.assertEqual(response.data["created_count"], 1 if attempt == 0 else 0)
            self.assertEqual(response.data["updated_count"], 1 if attempt == 0 else 2)
        self.assertEqual(Asset.objects.filter(tenant=self.tenant).count(), 2)
        self.asset.refresh_from_db()
        foreign.refresh_from_db()
        self.assertEqual(self.asset.custom_data, {"serial": "F-1", "capacity": 0, "active": False, "notes": "Keep this note"})
        self.assertEqual(self.asset.created_by_id, original_creator.pk)
        self.assertEqual(self.asset.assigned_to_id, self.user.pk)
        self.assertEqual(foreign.custom_data["capacity"], 99)
        self.assertEqual(Asset.objects.get(tenant=self.tenant, custom_data__serial="F-2").created_by_id, self.user.pk)

    def test_template_update_preserves_required_unmapped_fields(self):
        self.asset_type.schema[0]["is_unique_key"] = True
        self.asset_type.schema.append({"key": "notes", "label": "Notes", "type": "text", "required": True})
        self.asset_type.save()
        self.asset.custom_data["notes"] = "Existing required value"
        self.asset.save()
        response = self.client.post("/api/assets/items/bulk_import/", {
            "asset_type": str(self.asset_type.pk),
            "file": workbook_file([["Serial"], ["info"], ["example"], ["F-1"]]),
        }, format="multipart")
        self.assertEqual(response.data["error_count"], 0, response.data)
        self.assertEqual(response.data["updated_count"], 1)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.custom_data["notes"], "Existing required value")

    def test_template_import_without_unique_key_still_creates_records(self):
        response = self.client.post("/api/assets/items/bulk_import/", {
            "asset_type": str(self.asset_type.pk),
            "file": workbook_file([["Serial"], ["info"], ["example"], ["F-1"]]),
        }, format="multipart")
        self.assertEqual(response.data["created_count"], 1, response.data)
        self.assertEqual(response.data["updated_count"], 0)
        self.assertEqual(Asset.objects.count(), 2)

    def test_equipment_export_keeps_user_text_literal_in_excel(self):
        self.tenant.name = '=HYPERLINK("https://example.test", "yard")'
        self.tenant.save()
        self.asset_type.schema = [{"key": "notes", "label": "=2+2", "type": "text"}]
        self.asset_type.save()
        self.asset.custom_data = {"notes": '=HYPERLINK("https://example.test", "note")'}
        self.asset.save()
        response = self.client.get("/api/assets/items/export/", {"asset_type": str(self.asset_type.pk)})
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(BytesIO(response.content), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet["D1"].value, "=2+2")
        self.assertEqual(sheet["D1"].data_type, "s")
        self.assertEqual(sheet["B2"].value, self.tenant.name)
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["D2"].value, self.asset.custom_data["notes"])
        self.assertEqual(sheet["D2"].data_type, "s")
        workbook.close()

    def test_excel_downloads_normalize_invalid_and_reserved_sheet_names(self):
        for name in ("MIG/MAG kaynak makinesi", "_mapping", "_MaPpInG", "'[]:*?/\\'", "'" + "A" * 31 + "'", 'Vinç "özel"'):
            with self.subTest(name=name):
                self.asset_type.name = name
                self.asset_type.save()
                template = self.client.get(f"/api/assets/types/{self.asset_type.pk}/export_template/")
                self.assertEqual(template.status_code, 200)
                workbook = openpyxl.load_workbook(BytesIO(template.content))
                title = workbook.active.title
                self.assertTrue(title)
                self.assertLessEqual(len(title), 31)
                self.assertFalse(any(character in title for character in '[]:*?/\\'))
                self.assertFalse(title.startswith("'") or title.endswith("'"))
                self.assertNotEqual(title.casefold(), "_mapping")
                self.assertEqual(workbook["_mapping"].sheet_state, "hidden")
                disposition, parameters = parse_header_parameters(template["Content-Disposition"])
                self.assertEqual(disposition, "attachment")
                self.assertEqual(parameters["filename"], f"{title.replace(' ', '_')}_sablonu.xlsx")
                workbook.close()
                preview = self.client.post("/api/assets/items/preview_import/", {
                    "file": SimpleUploadedFile("equipment.xlsx", template.content),
                }, format="multipart")
                self.assertEqual(preview.status_code, 200, preview.data)
                self.assertEqual(preview.data["detected_template"]["asset_type_id"], str(self.asset_type.pk))
                exported = self.client.get("/api/assets/items/export/", {"asset_type": str(self.asset_type.pk)})
                self.assertEqual(exported.status_code, 200)
                exported_workbook = openpyxl.load_workbook(BytesIO(exported.content))
                self.assertEqual(exported_workbook.active.title, title)
                exported_workbook.close()
