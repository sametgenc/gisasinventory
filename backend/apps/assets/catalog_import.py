"""Preview-first, atomic Excel import for the global equipment type catalog."""
import hashlib
import json
from copy import deepcopy
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import iterparse
import re

import openpyxl
from django.core import signing
from django.db import IntegrityError, transaction
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response

from .catalog_constants import CATEGORIES, DEFAULT_SCHEMA
from .models import AssetType
from .serializers import AssetTypeSerializer
from apps.tenants.permissions import is_global_admin

HEADERS = {'Katalog kodu': 'catalog_code', 'Varlık tipi': 'name', 'Kategori kodu': 'category_code',
           'Ön paylaşım sınıfı': 'sharing_class', 'Kapsam dışı nedeni': 'exclusion_reason', 'Açıklama': 'description'}
MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
SALT = 'assets.catalog-import.v1'


def read_catalog(upload):
    if not upload or not upload.name.lower().endswith('.xlsx'):
        raise ValidationError({'detail': 'Lütfen .xlsx uzantılı katalog dosyası seçin.'})
    if upload.size > MAX_BYTES:
        raise ValidationError({'detail': 'Dosya en fazla 5 MB olabilir.'})
    content = upload.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValidationError({'detail': 'Dosya en fazla 5 MB olabilir.'})
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValidationError({'detail': 'Dosyanın açılmış boyutu çok büyük.'})
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = workbook['Katalog'] if 'Katalog' in workbook.sheetnames else workbook.active
            # Validate actual cells, not exporter-supplied dimensions: otherwise
            # data past column G could be silently omitted from the preview.
            with ZipFile(BytesIO(content)) as archive, archive.open(sheet._worksheet_path) as source:
                for _, element in iterparse(source, events=('end',)):
                    if element.tag.endswith('}c'):
                        coordinate = re.match(r'([A-Z]+)(\d+)', element.get('r', ''))
                        if coordinate and any(child.tag.endswith(('}v', '}is', '}f')) for child in element):
                            if coordinate.group(1) not in ('A', 'B', 'C', 'D', 'E', 'F'):
                                raise ValidationError({'detail': 'Şablonun altı sütunu dışında veri var. Ek sütunları kaldırın.'})
                    # Retain children until their enclosing cell has been inspected.
                    if element.tag.endswith('}c') or element.tag.endswith('}row'):
                        element.clear()
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(max_col=len(HEADERS) + 1)
            headers = [str(cell.value or '').strip() for cell in next(iterator, ())]
            while headers and not headers[-1]:
                headers.pop()
            mapping = [HEADERS.get(header, header) for header in headers]
            if len(mapping) != len(set(mapping)) or set(mapping) != set(HEADERS.values()):
                raise ValidationError({'detail': 'Sütunlar şablonla eşleşmiyor. Altı sütunlu katalog şablonunu kullanın.'})
            rows = []
            for row_number, cells in enumerate(iterator, 2):
                if row_number > MAX_ROWS + 1:
                    raise ValidationError({'detail': 'Tek dosyada en fazla 2000 satır aktarılabilir.'})
                if not any(cell.value is not None for cell in cells):
                    continue
                values = {key: str(cell.value).strip() if cell.value is not None else ''
                          for key, cell in zip(mapping, cells)}
                errors = []
                if cells[-1].value is not None:
                    errors.append('Şablonun dışında bir sütunda veri var.')
                if any(cell.data_type in ('f', 'e') for cell in cells):
                    errors.append('Formül veya Excel hatası içeren hücreleri düz metne dönüştürün.')
                for key in ('catalog_code', 'category_code', 'sharing_class'):
                    values[key] = values[key].upper()
                rows.append({'row': row_number, 'values': values, 'errors': errors})
        finally:
            workbook.close()
    except (BadZipFile, KeyError, ValueError, TypeError, OSError, SyntaxError):
        raise ValidationError({'detail': 'Excel dosyası okunamadı. Geçerli bir .xlsx dosyası yükleyin.'}) from None
    if not rows:
        raise ValidationError({'detail': 'Dosyada aktarılacak kayıt yok. Katalog sayfasına veri ekleyin.'})
    return content, rows


def catalog_revision():
    state = list(AssetType.objects.order_by('pk').values_list('pk', 'updated_at'))
    return hashlib.sha256(json.dumps(state, default=str).encode()).hexdigest()


def build_plan(rows, mode, request):
    existing = {item.catalog_code: item for item in AssetType.objects.all() if item.catalog_code}
    codes, names = set(), set()
    plans = []
    summary = {'create': 0, 'update': 0, 'skip': 0, 'error': 0}
    for row in rows:
        values, errors = row['values'], list(row['errors'])
        code, name = values['catalog_code'], values['name']
        if not code:
            errors.append('Katalog kodu zorunludur.')
        if not values['category_code']:
            errors.append('Kategori kodu zorunludur (A–M).')
        if not values['sharing_class']:
            errors.append('Ön paylaşım sınıfı zorunludur (A, B veya C).')
        if code in codes:
            errors.append('Katalog kodu dosyada birden fazla kez kullanılmış.')
        if name.casefold() in names:
            errors.append('Varlık tipi adı dosyada birden fazla kez kullanılmış.')
        codes.add(code)
        names.add(name.casefold())
        instance = existing.get(code)
        conflicting_name = AssetType.objects.filter(name__iexact=name)
        if instance:
            conflicting_name = conflicting_name.exclude(pk=instance.pk)
        if conflicting_name.exists():
            errors.append('Bu ad başka bir tipte kullanılıyor. Mevcut tipin katalog kodunu düzenleme ekranından eşleştirin.')
        serializer = AssetTypeSerializer(instance, data=values, partial=bool(instance), context={'request': request})
        if not serializer.is_valid():
            labels = {value: key for key, value in HEADERS.items()}
            errors.extend(f'{labels.get(key, key)}: {"; ".join(map(str, messages))}'
                          for key, messages in serializer.errors.items())
        operation = 'create'
        current_values = {key: getattr(instance, key) for key in values} if instance else None
        changes = {key: {'before': current_values[key], 'after': value}
                   for key, value in values.items() if current_values and current_values[key] != value}
        skip_reason = None
        if instance:
            unchanged = not changes
            operation = 'skip' if mode == 'create' or unchanged else 'update'
            if operation == 'skip':
                skip_reason = 'unchanged' if unchanged else 'exists'
        if errors:
            operation = 'error'
        summary[operation] += 1
        effective_values = current_values if operation == 'skip' else values
        plans.append({**row, 'errors': errors, 'operation': operation, 'serializer': serializer,
                      'current_values': current_values, 'changes': changes if operation == 'update' else {},
                      'skip_reason': skip_reason,
                      'sharing_excluded': effective_values['category_code'] == 'K' or effective_values['sharing_class'] == 'C' or bool(effective_values['exclusion_reason'])})
    return plans, summary


class CatalogImportMixin:
    def check_catalog_access(self, request):
        if not is_global_admin(request.user):
            raise PermissionDenied('Katalog aktarımı için platform yöneticisi yetkisi gerekir.')

    @action(detail=False, methods=['get'])
    def catalog_template(self, request):
        self.check_catalog_access(request)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Katalog'
        sheet.append(list(HEADERS))
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2563EB')
            sheet.column_dimensions[cell.column_letter].width = 28 if cell.column < 5 else 55
        for column, options in [('C', 'A,B,C,D,E,F,G,H,I,J,K,L,M'), ('D', 'A,B,C')]:
            validation = DataValidation(type='list', formula1=f'"{options}"', allow_blank=False)
            validation.errorTitle = 'Geçersiz kod'
            validation.error = 'Listeden bir kod seçin.'
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add(f'{column}2:{column}2001')
        guide = workbook.create_sheet('Rehber')
        for row in [
            ['KATALOG AKTARIMI'],
            ['Katalog sayfasına her varlık tipi için bir satır ekleyin. Başlıkları değiştirmeyin.'],
            ['Katalog kodu kalıcı kimliktir. Örn. B-001. Gerçek ekipman veya seri numarası değildir.'],
            ['Ön paylaşım sınıfı', 'A: Çok uygun / B: Koşullu uygun / C: Uygun değil'],
            ['Kategori K otomatik kapsam dışıdır. C sınıfı paylaşım değerlendirmesine uygun değildir.'],
            ['KKD, radyografi vb. istisnalar için kapsam dışı nedenini mutlaka doldurun.'],
            ['Ön sınıf paylaşım izni değildir; fiziksel ekipman için saha değerlendirmesi gerekir.'],
            ['Güncelleme katalog koduyla eşleşir; alan şeması ve ekipmanlar korunur.'],
            ['Yeni tiplerin alanları: Ekipman kodu (zorunlu, benzersiz), Marka, Model, Kapasite ve birimi.'],
            ['Boş hücreler güncelleme modunda mevcut açıklama/kapsam dışı nedenini temizler.'],
            ['Yalnızca Katalog sayfası aktarılır; aşağıdaki örnek aktarılmaz.'],
            ['Örnek kod', 'Örnek tip', 'Kategori', 'Ön sınıf'],
            ['B-001', 'Örnek kaldırma ekipmanı', 'B', 'B'],
            [], ['Kategori kodu', 'Kategori adı'], *CATEGORIES,
        ]:
            guide.append(row)
        guide.column_dimensions['A'].width = 105
        guide.column_dimensions['B'].width = 60
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="varlik-tipi-katalog-sablonu.xlsx"'
        return response

    def catalog_input(self, request):
        self.check_catalog_access(request)
        mode = request.data.get('mode', 'create')
        if mode not in ('create', 'update'):
            raise ValidationError({'detail': 'Geçersiz aktarım modu.'})
        content, rows = read_catalog(request.FILES.get('file'))
        identity = {'file': hashlib.sha256(content).hexdigest(), 'mode': mode, 'user': str(request.user.pk)}
        return rows, mode, identity

    @action(detail=False, methods=['post'])
    def catalog_preview(self, request):
        rows, mode, identity = self.catalog_input(request)
        revision = catalog_revision()
        plans, summary = build_plan(rows, mode, request)
        if revision != catalog_revision():
            return Response({'detail': 'Katalog değişti. Lütfen tekrar önizleyin.'}, status=409)
        token = signing.dumps({**identity, 'revision': revision}, salt=SALT) if not summary['error'] else None
        return Response({'rows': [{key: value for key, value in plan.items() if key != 'serializer'} for plan in plans],
                         'summary': summary, 'preview_token': token})

    @action(detail=False, methods=['post'])
    def catalog_import(self, request):
        rows, mode, identity = self.catalog_input(request)
        try:
            preview = signing.loads(request.data.get('preview_token', ''), salt=SALT, max_age=1800)
        except (signing.BadSignature, TypeError):
            raise ValidationError({'detail': 'Önizleme eksik veya süresi doldu. Dosyayı yeniden önizleyin.'}) from None
        if any(preview.get(key) != value for key, value in identity.items()):
            raise ValidationError({'detail': 'Dosya veya aktarım modu değişmiş. Yeniden önizleyin.'})
        try:
            with transaction.atomic():
                list(AssetType.objects.select_for_update().order_by('pk'))
                if preview['revision'] != catalog_revision():
                    return Response({'detail': 'Katalog önizlemeden sonra değişti. Yeniden önizleyin; hiçbir satır kaydedilmedi.'}, status=409)
                plans, summary = build_plan(rows, mode, request)
                if summary['error']:
                    return Response({'detail': 'Dosyada hatalar var. Yeniden önizleyin; hiçbir satır kaydedilmedi.'}, status=400)
                for plan in plans:
                    if plan['operation'] == 'create':
                        plan['serializer'].save(schema=deepcopy(DEFAULT_SCHEMA))
                    elif plan['operation'] == 'update':
                        plan['serializer'].save()
        except IntegrityError:
            return Response({'detail': 'Eşzamanlı bir değişiklik oluştu. Yeniden önizleyin; hiçbir satır kaydedilmedi.'}, status=409)
        return Response({'summary': summary})
