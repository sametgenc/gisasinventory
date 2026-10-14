"""Invented workbooks cover inventory suggestions and safe import review."""

import json
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APITestCase

from apps.tenants.models import Tenant
from apps.user.models import User
from .models import Asset, AssetType
from .validation import default_stock_schema


def inventory_file(rows, *, sheets=None, merge=None):
    workbook = openpyxl.Workbook()
    workbook.active.title = 'Envanter'
    for row in rows:
        workbook.active.append(row)
    for title, sheet_rows in sheets or []:
        sheet = workbook.create_sheet(title)
        for row in sheet_rows:
            sheet.append(row)
    for cell_range in merge or []:
        workbook.active.merge_cells(cell_range)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return SimpleUploadedFile('fictional-inventory.xlsx', stream.getvalue())


class InventoryWorkbookTests(APITestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(name='Invented inventory yard')
        self.other = Tenant.objects.create(name='Other invented yard')
        self.user = User.objects.create(username='inventory-importer', role='tenant_admin', tenant=self.tenant)
        self.stock = AssetType.objects.create(name='Örnek malzeme stoku', tracking_mode='stock',
                                             schema=default_stock_schema())
        self.client.force_authenticate(self.user)

    def preview(self, rows, **extra):
        return self.client.post('/api/assets/items/preview_import/', {
            'file': inventory_file(rows), **extra,
        }, format='multipart')

    def raw(self, rows, **extra):
        return self.client.post('/api/assets/items/raw_import/', {
            'file': inventory_file(rows), 'asset_type': str(self.stock.pk),
            'mapping': json.dumps({'Kod': 'material_code', 'İsim': 'material_name', 'Adet': 'quantity'}),
            **extra,
        }, format='multipart')

    def test_finds_header_below_title_and_skips_only_structural_rows(self):
        response = self.client.post('/api/assets/items/preview_import/', {
            'file': inventory_file([
                ['ÖRNEK ENVANTER LİSTESİ'], ['Hazırlanma tarihi', '2025-01-01'],
                ['Kod', 'İsim', 'Adet'], ['001', 'Örnek kelepçe', 0],
                ['GENEL TOPLAM', None, '=SUM(C4:C4)'],
                ['Kod', 'İsim', 'Adet'], ['İSKELE MALZEMELERİ'],
                ['002', 'Toplam kalite bağlantısı', 5], ['TOTAL', 'Örnek gerçek malzeme', 2],
            ], merge=['A1:C1', 'A7:C7']),
        }, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        analysis = response.data['analysis']
        self.assertEqual(analysis['header_row'], 3)
        self.assertEqual(analysis['suggested_asset_type_id'], str(self.stock.pk))
        self.assertEqual(analysis['counts'], {'data': 3, 'skipped': 3, 'review': 0, 'ready': 3})
        self.assertEqual(response.data['preview_row_numbers'], [4, 8, 9])
        self.assertEqual(response.data['preview_rows'][0], ['001', 'Örnek kelepçe', '0'])
        self.assertEqual([row['reason'] for row in analysis['skipped_rows']], ['total', 'repeated_header', 'section'])

    def test_multiple_equal_types_need_selection_but_explicit_type_is_authoritative(self):
        AssetType.objects.create(name='Another invented stock', tracking_mode='stock', schema=default_stock_schema())
        rows = [['Kod', 'İsim', 'Adet'], ['X-1', 'Clamp', 3]]
        response = self.preview(rows)
        self.assertIsNone(response.data['analysis']['suggested_asset_type_id'])
        self.assertIn('asset_type_selection_required', [item['code'] for item in response.data['analysis']['blockers']])
        chosen = self.preview(rows, asset_type=str(self.stock.pk))
        self.assertEqual(chosen.data['analysis']['analysis_asset_type_id'], str(self.stock.pk))
        self.assertFalse(chosen.data['analysis']['blockers'])

    def test_multiple_sheets_require_explicit_choice_when_equal(self):
        rows = [['Kod', 'İsim', 'Adet'], ['X-1', 'Clamp', 3]]
        response = self.client.post('/api/assets/items/preview_import/', {
            'file': inventory_file(rows, sheets=[('İkinci liste', rows)]),
        }, format='multipart')
        self.assertIn('sheet_selection_required', [item['code'] for item in response.data['analysis']['blockers']])
        selected = self.client.post('/api/assets/items/preview_import/', {
            'file': inventory_file(rows, sheets=[('İkinci liste', rows)]), 'sheet_name': 'İkinci liste',
        }, format='multipart')
        self.assertEqual(selected.data['analysis']['sheet_name'], 'İkinci liste')
        self.assertFalse(selected.data['analysis']['blockers'])

    def test_blank_header_can_be_manually_mapped_in_preview_and_import(self):
        rows = [['Kod', None, 'Adet'], ['001', 'Clamp', 0]]
        mapping = json.dumps({'Kod': 'material_code', 'Sütun 2': 'material_name', 'Adet': 'quantity'})
        preview = self.preview(rows, asset_type=str(self.stock.pk), header_row=1, mapping=mapping)
        self.assertEqual(preview.data['headers'], ['Kod', 'Sütun 2', 'Adet'])
        self.assertFalse(preview.data['analysis']['blockers'])
        response = self.raw(rows, header_row=1, mapping=mapping, require_review='true')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Asset.objects.get().custom_data['material_code'], '001')

    def test_trailing_blank_header_with_data_is_available_for_manual_mapping(self):
        rows = [['Kod', 'İsim', None], ['001', 'Clamp', 0]]
        mapping = json.dumps({'Kod': 'material_code', 'İsim': 'material_name', 'Sütun 3': 'quantity'})
        preview = self.preview(rows, asset_type=str(self.stock.pk), header_row=1, mapping=mapping)
        self.assertEqual(preview.data['headers'], ['Kod', 'İsim', 'Sütun 3'])
        self.assertEqual(preview.data['analysis']['counts']['ready'], 1)
        response = self.raw(rows, header_row=1, mapping=mapping, require_review='true')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Asset.objects.get().custom_data['quantity'], 0)

    def test_formula_error_and_invalid_quantity_need_review_before_any_writes(self):
        rows = [['Kod', 'İsim', 'Adet'], ['OK', 'Clamp', 3], ['F', 'Clamp', '=1+1'],
                ['E', 'Clamp', '#DIV/0!'], ['N', 'Clamp', -1]]
        preview = self.preview(rows, asset_type=str(self.stock.pk))
        self.assertEqual(preview.data['analysis']['review_row_numbers'], [3, 4, 5])
        self.assertEqual({item['reason'] for item in preview.data['analysis']['review_rows']},
                         {'formula', 'excel_error', 'invalid_value'})
        response = self.raw(rows, require_review='true')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Asset.objects.count(), 0)
        excluded = self.raw(rows, require_review='true', exclude_row_numbers='[3,4,5]')
        self.assertEqual(excluded.status_code, 200, excluded.data)
        self.assertEqual(excluded.data['created_count'], 1)
        self.assertEqual(excluded.data['skipped_count'], 3)

    def test_legacy_partial_import_still_reports_formula_rows_instead_of_zero(self):
        response = self.raw([['Kod', 'İsim', 'Adet'], ['OK', 'Clamp', 0], ['F', 'Clamp', '=1+1']])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual((response.data['created_count'], response.data['error_count']), (1, 1))
        self.assertEqual(response.data['errors'][0]['row'], 3)
        self.assertEqual(Asset.objects.get().custom_data['quantity'], 0)

    def test_duplicate_identifiers_both_need_review_and_are_not_summed(self):
        rows = [['Kod', 'İsim', 'Adet'], ['DUP', 'Clamp', 3], ['DUP', 'Clamp', 7]]
        response = self.preview(rows, asset_type=str(self.stock.pk))
        self.assertEqual(response.data['analysis']['review_row_numbers'], [2, 3])
        self.assertEqual(response.data['analysis']['counts']['ready'], 0)
        blocked = self.raw(rows, require_review='true')
        self.assertEqual(blocked.status_code, 400)
        self.assertEqual(Asset.objects.count(), 0)
        selected = self.raw(rows, require_review='true', exclude_row_numbers='[3]')
        self.assertEqual(selected.status_code, 200, selected.data)
        self.assertEqual(Asset.objects.get().custom_data['quantity'], 3)

    def test_same_model_fifty_times_is_valid_individual_inventory(self):
        equipment = AssetType.objects.create(name='Invented drills', schema=[
            {'key': 'serial', 'label': 'Serial number', 'type': 'text', 'required': True, 'is_unique_key': True},
            {'key': 'brand', 'label': 'Brand', 'type': 'text'},
            {'key': 'model', 'label': 'Model', 'type': 'text'},
        ])
        rows = [['Seri No', 'Marka', 'Model']] + [[f'D-{number}', 'Fictional', 'Common model'] for number in range(50)]
        response = self.preview(rows, asset_type=str(equipment.pk))
        self.assertEqual(response.data['analysis']['counts']['ready'], 50)
        self.assertFalse(response.data['analysis']['review_rows'])

    def test_duplicate_headers_and_conflicting_suggestions_cannot_silently_choose_column(self):
        for headers, code in [(['Kod', 'Kod', 'İsim', 'Adet'], 'duplicate_headers'),
                              (['Kod', 'Malzeme kodu', 'İsim', 'Adet'], 'ambiguous_mapping')]:
            with self.subTest(headers=headers):
                response = self.preview([headers, ['A', 'B', 'Clamp', 3]], asset_type=str(self.stock.pk), header_row=1)
                self.assertIn(code, [item['code'] for item in response.data['analysis']['blockers']])
                self.assertIsNone(response.data['analysis']['suggested_mapping']['Kod'])

    def test_raw_and_template_paths_skip_same_totals_without_owner_inference(self):
        foreign = Asset.objects.create(tenant=self.other, asset_type=self.stock,
                                       custom_data={'material_code': '001', 'material_name': 'Foreign', 'quantity': 99})
        for template in [False, True]:
            rows = [['Malzeme kodu', 'Malzeme adı', 'Stok adedi']]
            if template:
                rows += [['info'], ['example']]
            rows += [['001', 'Clamp', 0], ['TOPLAM', None, 0]]
            response = self.client.post(f'/api/assets/items/{"bulk_import" if template else "raw_import"}/', {
                'asset_type': str(self.stock.pk), 'file': inventory_file(rows), 'tenant': self.other.pk,
                'mapping': json.dumps({'Malzeme kodu': 'material_code', 'Malzeme adı': 'material_name', 'Stok adedi': 'quantity'}),
            }, format='multipart')
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data['skipped_count'], 1)
            self.assertEqual(response.data['created_count'] + response.data['updated_count'], 1)
        foreign.refresh_from_db()
        self.assertEqual(foreign.custom_data['quantity'], 99)
        self.assertEqual(Asset.objects.get(tenant=self.tenant).custom_data['quantity'], 0)

    def test_caps_preview_details_without_losing_counts_or_review_row_numbers(self):
        rows = [['Kod', 'İsim', 'Adet']] + [[f'X-{number}', 'Clamp', '=1+1'] for number in range(70)]
        rows += [['TOPLAM', None, 1]] * 120
        response = self.preview(rows, asset_type=str(self.stock.pk))
        analysis = response.data['analysis']
        self.assertEqual(analysis['counts'], {'data': 70, 'skipped': 120, 'review': 70, 'ready': 0})
        self.assertEqual(len(analysis['skipped_rows']), 100)
        self.assertEqual(len(analysis['review_rows']), 50)
        self.assertEqual(len(analysis['review_row_numbers']), 70)
        self.assertEqual(len(response.data['preview_rows']), 50)

    def test_invalid_layout_exclusions_and_bounded_files_are_explicit_errors(self):
        rows = [['Kod', 'İsim', 'Adet'], ['X', 'Clamp', 1]]
        for extra in [{'header_row': '0'}, {'header_row': 'bad'}, {'sheet_name': 'Missing'},
                      {'exclude_row_numbers': '[true]'}, {'exclude_row_numbers': '[-2]'},
                      {'data_start_row': '1'}]:
            with self.subTest(extra=extra):
                response = self.raw(rows, **extra)
                self.assertEqual(response.status_code, 400, response.data)
        with patch('apps.assets.inventory_import.MAX_PHYSICAL_ROWS', 2):
            response = self.preview(rows + [['Y', 'Clamp', 2]])
            self.assertEqual(response.status_code, 400, response.data)
        with patch('apps.assets.inventory_import.MAX_UPLOAD_BYTES', 10):
            response = self.preview(rows)
            self.assertEqual(response.status_code, 400, response.data)

    def test_unmerged_uppercase_record_is_not_discarded_as_section_heading(self):
        equipment = AssetType.objects.create(name='Simple invented records', schema=[
            {'key': 'name', 'label': 'Name', 'type': 'text', 'required': True},
        ])
        response = self.preview([['Name'], ['GENERATOR']], asset_type=str(equipment.pk))
        self.assertEqual(response.data['analysis']['counts']['ready'], 1)
        self.assertFalse(response.data['analysis']['skipped_rows'])

    def test_single_column_total_named_record_is_preserved_without_total_structure(self):
        equipment = AssetType.objects.create(name='Single column names', schema=[
            {'key': 'name', 'label': 'Name', 'type': 'text', 'required': True},
        ])
        response = self.preview([['Name'], ['TOTAL']], asset_type=str(equipment.pk))
        self.assertEqual(response.data['analysis']['counts']['ready'], 1)
        self.assertEqual(response.data['preview_rows'], [['TOTAL']])

    def test_total_named_record_with_blank_optional_brand_is_preserved(self):
        equipment = AssetType.objects.create(name='Named individual equipment', schema=[
            {'key': 'name', 'label': 'Name', 'type': 'text', 'required': True},
            {'key': 'brand', 'label': 'Brand', 'type': 'text'},
        ])
        rows = [['Name', 'Brand'], ['TOTAL', None]]
        preview = self.preview(rows, asset_type=str(equipment.pk))
        self.assertEqual(preview.data['analysis']['counts']['ready'], 1)
        self.assertEqual(preview.data['analysis']['counts']['skipped'], 0)
        response = self.raw(rows, asset_type=str(equipment.pk), require_review='true',
                            mapping='{"Name":"name","Brand":"brand"}')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Asset.objects.get().custom_data['name'], 'TOTAL')

    def test_total_label_with_nonaggregate_formula_still_needs_manual_review(self):
        response = self.preview([['Kod', 'İsim', 'Adet'], ['TOTAL', None, '=1+1']], asset_type=str(self.stock.pk))
        self.assertEqual(response.data['analysis']['counts']['skipped'], 0)
        self.assertEqual(response.data['analysis']['review_row_numbers'], [2])
        self.assertIn('formula', [row['reason'] for row in response.data['analysis']['review_rows']])

    def test_numeric_identifier_and_exact_total_material_name_are_real_record(self):
        rows = [['Kod', 'İsim', 'Adet'], [0, 'TOTAL', 3]]
        preview = self.preview(rows)
        self.assertEqual(preview.data['analysis']['counts']['ready'], 1)
        response = self.raw(rows, require_review='true')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Asset.objects.get().custom_data, {
            'material_code': '0', 'material_name': 'TOTAL', 'quantity': 3,
        })

    def test_mixed_types_require_separate_import_and_exact_single_type_resolves_tie(self):
        other_type = AssetType.objects.create(name='İkinci örnek stok', tracking_mode='stock', schema=default_stock_schema())
        headers = ['Kod', 'İsim', 'Adet', 'Ekipman türü']
        rows = [headers, ['A', 'Clamp', 3, self.stock.name], ['B', 'Clamp', 4, other_type.name]]
        mixed = self.preview(rows, asset_type=str(self.stock.pk))
        self.assertIn('mixed_asset_types', [item['code'] for item in mixed.data['analysis']['blockers']])
        blocked = self.raw(rows, require_review='true')
        self.assertEqual(blocked.status_code, 400, blocked.data)
        self.assertEqual(Asset.objects.count(), 0)
        selected = self.preview(rows[:2])
        self.assertEqual(selected.data['analysis']['suggested_asset_type_id'], str(self.stock.pk))
        self.assertFalse(selected.data['analysis']['blockers'])

    def test_manual_skip_rules_are_explicit_and_preview_matches_import(self):
        rows = [['Kod', 'İsim', 'Adet'], ['SKIP', 'Explicit skipped record', 10], ['TOTAL', 'Real record', 0]]
        preview = self.preview(rows, asset_type=str(self.stock.pk), skip_rows='["SKIP"]')
        self.assertEqual(preview.data['preview_row_numbers'], [3])
        self.assertEqual(preview.data['analysis']['skipped_rows'][0]['reason'], 'custom')
        response = self.raw(rows, require_review='true', skip_rows='["SKIP"]')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['skipped_count'], 1)
        self.assertEqual(Asset.objects.get().custom_data['material_code'], 'TOTAL')

    def test_unknown_metadata_is_warning_only_and_does_not_infer_owner(self):
        rows = [['Kod', 'İsim', 'Adet', 'Tersane', 'Kaynak ID'], ['X', 'Clamp', 0, 'Other yard', 'source-example']]
        preview = self.preview(rows)
        self.assertEqual(preview.data['analysis']['counts']['ready'], 1)
        self.assertFalse(preview.data['analysis']['review_rows'])
        self.assertEqual(preview.data['analysis']['warnings'][0]['code'], 'unmatched_columns')
        response = self.raw(rows, require_review='true')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Asset.objects.get().tenant_id, self.tenant.pk)

    def test_formula_looking_literal_text_is_not_executed_or_mistaken_for_formula(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(['Kod', 'İsim', 'Adet'])
        workbook.active.append(['=INVENTED-CODE', 'Clamp', 0])
        workbook.active['A2'].data_type = 's'
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        response = self.client.post('/api/assets/items/raw_import/', {
            'asset_type': str(self.stock.pk), 'require_review': 'true',
            'mapping': json.dumps({'Kod': 'material_code', 'İsim': 'material_name', 'Adet': 'quantity'}),
            'file': SimpleUploadedFile('literal.xlsx', stream.getvalue()),
        }, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Asset.objects.get().custom_data['material_code'], '=INVENTED-CODE')

    def test_header_ambiguity_requires_manual_selection_and_explicit_layout_wins(self):
        rows = [['Kod', 'İsim', 'Adet'], ['A', 'Clamp', 0], ['Malzeme kodu', 'Malzeme adı', 'Miktar'], ['B', 'Clamp', 2]]
        automatic = self.preview(rows, asset_type=str(self.stock.pk))
        self.assertIn('header_selection_required', [item['code'] for item in automatic.data['analysis']['blockers']])
        selected = self.preview(rows, asset_type=str(self.stock.pk), header_row=3)
        self.assertFalse(selected.data['analysis']['blockers'])
        self.assertEqual(selected.data['preview_row_numbers'], [4])

    def test_expanded_size_and_column_limits_reject_instead_of_truncating(self):
        rows = [['Kod', 'İsim', 'Adet'], ['A', 'Clamp', 0]]
        with patch('apps.assets.inventory_import.MAX_EXPANDED_BYTES', 100):
            self.assertEqual(self.preview(rows).status_code, 400)
        with patch('apps.assets.inventory_import.MAX_COLUMNS', 2):
            self.assertEqual(self.preview(rows).status_code, 400)

    def test_sparse_sheet_grid_budget_rejects_before_workbook_allocation(self):
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = 'Kod'
        workbook.active['IV20000'] = 'Invented far-away cell'
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        with patch('apps.assets.inventory_import.openpyxl.load_workbook') as loader:
            response = self.client.post('/api/assets/items/preview_import/', {
                'file': SimpleUploadedFile('sparse.xlsx', stream.getvalue()),
            }, format='multipart')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('1 milyon', response.data['detail'])
        loader.assert_not_called()

    def test_aliased_worksheet_relationships_reject_before_any_cell_iteration(self):
        # One three-cell XML grid linked twice exceeds the patched five-cell
        # allocation budget, although the archive stores the grid only once.
        source = inventory_file([['Kod', 'İsim', 'Adet']], sheets=[('Alias', [['Placeholder']])])
        stream = BytesIO()
        with ZipFile(BytesIO(source.read())) as archive, ZipFile(stream, 'w') as rewritten:
            for name in archive.namelist():
                if name == 'xl/worksheets/sheet2.xml':
                    continue
                value = archive.read(name)
                if name == 'xl/_rels/workbook.xml.rels':
                    value = value.replace(b'/xl/worksheets/sheet2.xml', b'/xl/worksheets/sheet1.xml')
                rewritten.writestr(name, value)
        with patch('apps.assets.inventory_import.MAX_ANALYSIS_CELLS', 5), patch(
                'openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows') as iterator:
            response = self.client.post('/api/assets/items/preview_import/', {
                'file': SimpleUploadedFile('alias.xlsx', stream.getvalue()),
            }, format='multipart')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('yinelenen bağlantılar', response.data['detail'])
        iterator.assert_not_called()

    def test_exaggerated_dimension_uses_actual_extents_without_padding(self):
        source = inventory_file([['Kod', 'İsim', 'Adet'], ['X', 'Clamp', 0]])
        stream = BytesIO()
        with ZipFile(BytesIO(source.read())) as archive, ZipFile(stream, 'w') as rewritten:
            for name in archive.namelist():
                value = archive.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    value = value.replace(b'ref="A1:C2"', b'ref="A1:IV20000"')
                rewritten.writestr(name, value)
        response = self.client.post('/api/assets/items/preview_import/', {
            'file': SimpleUploadedFile('dimension.xlsx', stream.getvalue()),
        }, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['headers'], ['Kod', 'İsim', 'Adet'])
        self.assertEqual(response.data['analysis']['counts']['data'], 1)

    def test_template_selected_header_mapping_and_skips_match_preview(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(['Previous title'])
        workbook.active.append(['Previous info'])
        workbook.active.append(['Kod', 'İsim', 'Adet'])
        workbook.active.append(['SKIP', 'Explicitly omitted', 4])
        workbook.active.append(['X', 'Clamp', 0])
        workbook.active.append(['TOPLAM', None, 0])
        metadata = workbook.create_sheet('_mapping')
        metadata.sheet_state = 'hidden'
        metadata['A2'] = str(self.stock.pk)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        options = {
            'asset_type': str(self.stock.pk), 'header_row': 3, 'skip_rows': '["SKIP"]',
            'mapping': json.dumps({'Kod': 'material_code', 'İsim': 'material_name', 'Adet': 'quantity'}),
        }
        preview = self.client.post('/api/assets/items/preview_import/', {
            **options, 'file': SimpleUploadedFile('template.xlsx', stream.getvalue()),
        }, format='multipart')
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data['analysis']['data_start_row'], 4)
        self.assertEqual(preview.data['preview_row_numbers'], [5])
        response = self.client.post('/api/assets/items/bulk_import/', {
            **options, 'require_review': 'true', 'file': SimpleUploadedFile('template.xlsx', stream.getvalue()),
        }, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual((response.data['success_count'], response.data['skipped_count']), (1, 2))
        self.assertEqual(Asset.objects.get().custom_data['material_code'], 'X')

    def test_template_strict_review_blocks_before_valid_row_write(self):
        rows = [['Malzeme kodu', 'Malzeme adı', 'Stok adedi'], ['info'], ['example'],
                ['A', 'Clamp', 0], ['B', 'Clamp', '=2+2']]
        response = self.client.post('/api/assets/items/bulk_import/', {
            'asset_type': str(self.stock.pk), 'file': inventory_file(rows), 'require_review': 'true',
        }, format='multipart')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Asset.objects.count(), 0)
        self.assertEqual(response.data['analysis']['review_row_numbers'], [5])
