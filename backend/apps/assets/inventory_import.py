"""Bounded, deterministic inventory analysis shared by preview and both imports.

Workbook text is data. Formulas are never evaluated and uncertain rows require
review; existing callers can still use the established per-row import results.
"""

from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from io import BytesIO
import json
import re
import unicodedata
from typing import Protocol, TypedDict, cast
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import openpyxl
from openpyxl.workbook.workbook import Workbook
from rest_framework.exceptions import ValidationError

from .models import AssetType
from .validation import JSONValue, import_field_value, validate_custom_data


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 30 * 1024 * 1024
MAX_PHYSICAL_ROWS = 20_000
MAX_COLUMNS = 256
MAX_SHEETS = 32
MAX_ARCHIVE_MEMBERS = 512
MAX_ANALYSIS_CELLS = 1_000_000
HEADER_SCAN_ROWS = 50
TOTAL_LABELS = {'toplam', 'geneltoplam', 'aratoplam', 'total', 'subtotal', 'grandtotal'}
TOTAL_FORMULA = re.compile(r'^\s*=\s*(?:SUM|SUBTOTAL|AGGREGATE|SUMIF|SUMIFS)\s*\(', re.IGNORECASE)
TYPE_HEADERS = {'ekipmanturu', 'varliktipi', 'assettype', 'equipmenttype'}
ALIASES = {
    'code': {'materialcode', 'stockcode', 'malzemekodu', 'stokkodu', 'urun kodu', 'kod', 'code'},
    'serial': {'serial', 'serialnumber', 'serialno', 'serino', 'serinumarasi', 'serinumarası'},
    'name': {'materialname', 'material_name', 'malzemeadi', 'stokadi', 'urunadi', 'isim', 'adi', 'name'},
    'quantity': {'quantity', 'qty', 'count', 'adet', 'miktar', 'stokadedi', 'stokmiktari'},
    'brand': {'brand', 'manufacturer', 'marka', 'uretici'},
    'model': {'model', 'modelname', 'modeladi'},
}


class InventoryImportError(ValueError):
    """Invalid input or a workbook exceeding an explicit resource limit."""


class WorkbookSource(Protocol):
    def seek(self, offset: int) -> int: ...

    def read(self, size: int = -1) -> bytes: ...


class ImportOptions(TypedDict):
    mapping: dict[str, str | None] | None
    header_row: int | None
    data_start_row: int | None
    sheet_name: str | None
    skip_rows: list[str]
    exclude_row_numbers: set[int]
    require_review: bool


class SheetMetadata(TypedDict):
    merged_headings: set[int]
    max_row: int
    max_column: int


def normalized(value: object) -> str:
    text = str(value or '').strip().replace('ı', 'i').replace('İ', 'I').casefold()
    text = unicodedata.normalize('NFKD', text)
    return ''.join(character for character in text if character.isalnum() and not unicodedata.combining(character))


ALIASES = {group: {normalized(alias) for alias in aliases} for group, aliases in ALIASES.items()}


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def json_parameter(value: object, default: object, label: str) -> object:
    if value is None:
        return default
    try:
        return json.loads(value) if isinstance(value, str) else value
    except (json.JSONDecodeError, TypeError) as exc:
        raise InventoryImportError(f'{label} geçerli JSON olmalıdır.') from exc


def positive_integer(value: object, label: str) -> int | None:
    if value in (None, ''):
        return None
    if isinstance(value, bool) or not re.fullmatch(r'[0-9]+', str(value)):
        raise InventoryImportError(f'{label} pozitif tam sayı olmalıdır.')
    result = int(value)
    if result < 1 or result > MAX_PHYSICAL_ROWS:
        raise InventoryImportError(f'{label} 1 ile {MAX_PHYSICAL_ROWS} arasında olmalıdır.')
    return result


def import_options(data: Mapping[str, object]) -> ImportOptions:
    mapping = json_parameter(data.get('mapping'), None, 'Sütun eşleştirmesi')
    if mapping is not None and (not isinstance(mapping, dict) or any(
            not isinstance(header, str) or (key is not None and not isinstance(key, str))
            for header, key in mapping.items())):
        raise InventoryImportError('Sütun eşleştirmesi bir alan listesi olmalıdır.')
    skip_rows = json_parameter(data.get('skip_rows'), [], 'Atlanacak etiketler')
    if not isinstance(skip_rows, list) or not all(isinstance(item, str) for item in skip_rows):
        raise InventoryImportError('Atlanacak etiketler metin listesi olmalıdır.')
    exclusions = json_parameter(data.get('exclude_row_numbers'), [], 'Atlanacak satırlar')
    if (not isinstance(exclusions, list) or any(type(number) is not int or number < 1
                                               or number > MAX_PHYSICAL_ROWS for number in exclusions)):
        raise InventoryImportError('Atlanacak satırlar pozitif tam sayı listesi olmalıdır.')
    sheet_name = data.get('sheet_name')
    if sheet_name is not None and (not isinstance(sheet_name, str) or not sheet_name.strip()):
        raise InventoryImportError('Excel sayfası geçerli bir ad olmalıdır.')
    require_review = data.get('require_review', False)
    if require_review not in (False, True, 'false', 'true', 'False', 'True', '0', '1'):
        raise InventoryImportError('İnceleme tercihi geçersiz.')
    return {
        'mapping': cast(dict[str, str | None] | None, mapping),
        'header_row': positive_integer(data.get('header_row'), 'Başlık satırı'),
        'data_start_row': positive_integer(data.get('data_start_row'), 'İlk veri satırı'),
        'sheet_name': cast(str | None, sheet_name), 'skip_rows': cast(list[str], skip_rows),
        'exclude_row_numbers': set(cast(list[int], exclusions)),
        'require_review': require_review in (True, 'true', 'True', '1'),
    }


@dataclass(frozen=True)
class Cell:
    value: object = None
    kind: str = 'n'


@dataclass
class Sheet:
    name: str
    rows: list[tuple[Cell, ...]]
    merged_headings: set[int]


@dataclass
class InventoryAnalysis:
    preview: dict[str, object]
    rows: list[tuple[int, tuple[Cell, ...]]]
    converted: dict[int, dict[str, JSONValue]]
    row_errors: dict[int, list[str]]
    skipped_count: int

    @property
    def analysis(self) -> dict[str, object]:
        return cast(dict[str, object], self.preview['analysis'])


def _column_number(reference: str | None) -> int:
    match = re.match(r'([A-Z]+)', reference or '')
    if not match:
        return 0
    number = 0
    for letter in match.group(1):
        number = number * 26 + ord(letter) - 64
        if number > MAX_COLUMNS:
            return number
    return number


def _bounded_workbook(file: WorkbookSource) -> tuple[Workbook, dict[str, SheetMetadata]]:
    file.seek(0)
    raw = file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise InventoryImportError('Excel dosyası en fazla 5 MB olabilir.')
    metadata_by_path = {}
    try:
        with ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if (len(members) > MAX_ARCHIVE_MEMBERS
                    or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES):
                raise InventoryImportError('Excel içeriği çok büyük: açılmış dosya sınırı 30 MB.')
            if len({member.filename for member in members}) != len(members):
                raise InventoryImportError('Excel arşivinde yinelenen dosyalar var.')
            aggregate_cells = 0
            for member in members:
                if not re.fullmatch(r'xl/worksheets/[^/]+\.xml', member.filename):
                    continue
                headings = set()
                physical_rows = 0
                max_row = 0
                max_column = 0
                current_row = 0
                current_column = 0
                with archive.open(member) as stream:
                    for event, element in ElementTree.iterparse(stream, events=('start', 'end')):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if event == 'start' and tag == 'row':
                            physical_rows += 1
                            row_number = int(element.attrib.get('r', max_row + 1))
                            if physical_rows > MAX_PHYSICAL_ROWS or not 1 <= row_number <= MAX_PHYSICAL_ROWS:
                                raise InventoryImportError(f'Excel sayfası en fazla {MAX_PHYSICAL_ROWS} satır olabilir.')
                            if row_number <= max_row:
                                raise InventoryImportError('Excel satır sırası geçersiz.')
                            max_row = max(max_row, row_number)
                            current_row = row_number
                            current_column = 0
                        elif event == 'start':
                            continue
                        elif tag == 'row':
                            element.clear()
                        elif tag == 'c':
                            reference = element.attrib.get('r')
                            match = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', reference) if reference else None
                            if reference and (match is None or int(match.group(2)) != current_row):
                                raise InventoryImportError('Excel hücre konumu geçersiz.')
                            column_number = _column_number(reference) if reference else current_column + 1
                            if column_number <= current_column:
                                raise InventoryImportError('Excel sütun sırası geçersiz.')
                            if column_number > MAX_COLUMNS:
                                raise InventoryImportError(f'Excel sayfası en fazla {MAX_COLUMNS} sütun olabilir.')
                            max_column = max(max_column, column_number)
                            current_column = column_number
                            element.clear()
                        elif tag == 'dimension':
                            end = element.attrib.get('ref', '').split(':')[-1]
                            number = re.search(r'(\d+)$', end)
                            if ((number and int(number.group(1)) > MAX_PHYSICAL_ROWS)
                                    or _column_number(end) > MAX_COLUMNS):
                                raise InventoryImportError('Excel sayfa boyutları izin verilen sınırı aşıyor.')
                            element.clear()
                        elif tag == 'mergeCell':
                            bounds = openpyxl.utils.cell.range_boundaries(element.attrib['ref'])
                            if bounds[2] > MAX_COLUMNS or bounds[3] > MAX_PHYSICAL_ROWS:
                                raise InventoryImportError('Excel birleştirilmiş alanı izin verilen sınırı aşıyor.')
                            if bounds[1] == bounds[3] and bounds[2] > bounds[0]:
                                headings.add(bounds[1])
                            element.clear()
                aggregate_cells += max_row * max_column
                if aggregate_cells > MAX_ANALYSIS_CELLS:
                    raise InventoryImportError('Excel inceleme alanı çok büyük: sayfalar toplamda en fazla 1 milyon hücre içerebilir. Dosyayı daha küçük listelere ayırın.')
                metadata_by_path[member.filename] = {
                    'merged_headings': headings, 'max_row': max_row, 'max_column': max_column,
                }
        workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except InventoryImportError:
        raise
    except (BadZipFile, ElementTree.ParseError, KeyError, ValueError, OSError) as exc:
        raise InventoryImportError('Dosya okunamadı. Geçerli bir .xlsx dosyası seçin.') from exc
    return workbook, metadata_by_path


def _validate_worksheet_relations(workbook: Workbook, metadata_by_path: dict[str, SheetMetadata]) -> None:
    if len(workbook.worksheets) > MAX_SHEETS:
        raise InventoryImportError(f'Excel dosyasında en fazla {MAX_SHEETS} sayfa olabilir.')
    paths = [worksheet._worksheet_path.lstrip('/') for worksheet in workbook.worksheets]
    if len(paths) != len(set(paths)):
        raise InventoryImportError('Excel sayfaları aynı içeriğe yinelenen bağlantılar içeriyor.')
    if any(path not in metadata_by_path for path in paths):
        raise InventoryImportError('Excel sayfa yapısı okunamadı.')


def _read_sheets(workbook: Workbook, metadata_by_path: dict[str, SheetMetadata]) -> list[Sheet]:
    sheets = []
    for worksheet in workbook.worksheets:
        if worksheet.title == '_mapping' or worksheet.sheet_state != 'visible':
            continue
        path = worksheet._worksheet_path.lstrip('/')
        metadata = metadata_by_path.get(path)
        if metadata is None:
            raise InventoryImportError('Excel sayfa yapısı okunamadı.')
        rows = []
        # Worksheet dimension metadata can exaggerate a sparse file's size.
        # Iterate only XML-verified extents, after the aggregate allocation cap.
        for number, row in enumerate(worksheet.iter_rows(
                min_row=1, min_col=1, max_row=metadata['max_row'] or 1,
                max_col=metadata['max_column'] or 1), start=1):
            if number > MAX_PHYSICAL_ROWS or len(row) > MAX_COLUMNS:
                raise InventoryImportError('Excel sayfa boyutları izin verilen sınırı aşıyor.')
            rows.append(tuple(Cell(cell.value, cell.data_type) for cell in row))
        sheets.append(Sheet(worksheet.title, rows, metadata['merged_headings']))
    if not sheets:
        raise InventoryImportError('Dosyada görünür bir envanter sayfası bulunamadı.')
    return sheets


def _field_aliases(field, tracking_mode):
    names = {normalized(field.get('key', '')), normalized(field.get('label', ''))}
    groups = {group for group, aliases in ALIASES.items() if names.intersection(aliases)}
    if tracking_mode == 'stock' and field.get('key') == 'material_code':
        groups.update(('code', 'serial'))
    for group in groups:
        names.update(ALIASES[group])
    return names - {''}


def _suggest_mapping(headers, asset_type):
    schema = asset_type.schema or []
    candidates = {}
    duplicates = {header for header, count in Counter(headers).items() if count > 1}
    for header in headers:
        matches = [field['key'] for field in schema
                   if normalized(header) in _field_aliases(field, asset_type.tracking_mode)]
        candidates[header] = matches if header not in duplicates else []
    target_counts = Counter(key for keys in candidates.values() for key in keys if len(keys) == 1)
    mapping = {header: keys[0] if len(keys) == 1 and target_counts[keys[0]] == 1 else None
               for header, keys in candidates.items()}
    ambiguous = [header for header, keys in candidates.items()
                 if len(keys) > 1 or (len(keys) == 1 and target_counts[keys[0]] > 1)]
    return mapping, ambiguous


def _headers(row, column_count=None):
    values = [cell.value for cell in row]
    if column_count is None:
        while values and is_blank(values[-1]):
            values.pop()
    else:
        values = (values + [None] * column_count)[:column_count]
    return [str(value).strip() if not is_blank(value) else f'Sütun {index}'
            for index, value in enumerate(values, start=1)]


def _header_candidates(sheet, asset_types):
    candidates = []
    seen = set()
    for number, row in enumerate(sheet.rows[:HEADER_SCAN_ROWS], start=1):
        if number in sheet.merged_headings:
            continue
        filled = [cell for cell in row if not is_blank(cell.value)]
        if not filled or any(cell.kind in ('f', 'e') or not isinstance(cell.value, str) for cell in filled):
            continue
        headers = _headers(row)
        identity = tuple(normalized(header) for header in headers)
        if identity in seen:
            continue  # Repeated table headers do not create layout ambiguity.
        seen.add(identity)
        matches = max((sum(key is not None for key in _suggest_mapping(headers, asset_type)[0].values())
                       for asset_type in asset_types), default=0)
        known = sum(any(normalized(header) in aliases for aliases in ALIASES.values()) for header in headers)
        score = (matches, known)
        if matches or known >= 2:
            candidates.append((score, number, headers))
    return sorted(candidates, key=lambda candidate: (candidate[0], -candidate[1]), reverse=True)


def _layout(sheets, asset_types, sheet_name, header_row, template):
    blockers = []
    sheet_candidates = [(sheet, _header_candidates(sheet, asset_types)) for sheet in sheets]
    if sheet_name is not None:
        selected = next((item for item in sheet_candidates if item[0].name == sheet_name), None)
        if selected is None:
            raise InventoryImportError('Seçilen Excel sayfası bulunamadı.')
    else:
        ranked = sorted(sheet_candidates, key=lambda item: item[1][0][0] if item[1] else (-1, -1), reverse=True)
        selected = ranked[0]
        if len(ranked) > 1:
            best_score = selected[1][0][0] if selected[1] else (-1, -1)
            ties = [item for item in ranked if (item[1][0][0] if item[1] else (-1, -1)) == best_score]
            if len(ties) > 1:
                blockers.append({'code': 'sheet_selection_required'})
    sheet, candidates = selected
    if header_row is not None:
        if header_row > len(sheet.rows):
            raise InventoryImportError('Başlık satırı seçilen sayfada bulunamadı.')
        chosen_row = header_row
    elif template:
        chosen_row = 1
    elif candidates:
        chosen_row = candidates[0][1]
        if sum(candidate[0] == candidates[0][0] for candidate in candidates) > 1:
            blockers.append({'code': 'header_selection_required'})
    else:
        chosen_row = next((number for number, row in enumerate(sheet.rows, start=1)
                           if number not in sheet.merged_headings and any(not is_blank(cell.value) for cell in row)), 1)
        blockers.append({'code': 'header_selection_required'})
    used_columns = max((index + 1 for row in sheet.rows[chosen_row - 1:]
                        for index, cell in enumerate(row) if not is_blank(cell.value)), default=0)
    headers = _headers(sheet.rows[chosen_row - 1], used_columns) if sheet.rows else []
    if not headers and not any(item['code'] == 'header_selection_required' for item in blockers):
        blockers.append({'code': 'header_selection_required'})
    return sheet, chosen_row, headers, blockers


def _row_skip_reason(number, cells, headers, sheet, skip_rows, mapping, asset_type):
    values = [cell.value for cell in cells]
    present = [(index, cell) for index, cell in enumerate(cells) if not is_blank(cell.value)]
    if not present:
        return 'blank'
    first = present[0][1].value
    if isinstance(first, str) and first.strip().casefold() in skip_rows:
        return 'custom'
    if len(values) >= len(headers) and all(normalized(value) == normalized(header)
                                          if not header.startswith('Sütun ') else is_blank(value)
                                          for value, header in zip(values, headers)):
        return 'repeated_header'
    if number in sheet.merged_headings and len(present) == 1 and isinstance(first, str):
        return 'section'
    total_cells = [(index, cell) for index, cell in present
                   if isinstance(cell.value, str) and normalized(cell.value) in TOTAL_LABELS and cell.kind != 'f']
    if len(total_cells) == 1 and len(headers) > 1:
        unique_fields = {field['key'] for field in asset_type.schema or [] if field.get('is_unique_key')} if asset_type else set()
        code_columns = [index for index, header in enumerate(headers)
                        if normalized(header) in ALIASES['code'] | ALIASES['serial']
                        or normalized(mapping.get(header)) in ALIASES['code'] | ALIASES['serial']
                        or mapping.get(header) in unique_fields]
        name_columns = [index for index, header in enumerate(headers)
                        if normalized(header) in ALIASES['name'] or normalized(mapping.get(header)) in ALIASES['name']]
        if (any(index < len(cells) and not is_blank(cells[index].value) for index in code_columns)
                and any(index < len(cells) and not is_blank(cells[index].value) for index in name_columns)):
            return None
        total_index = total_cells[0][0]
        others = [cell for index, cell in present if index != total_index]
        # A real code and material name give two text cells, so "TOTAL" can be a real code.
        if others and all((cell.kind == 'f' and TOTAL_FORMULA.match(str(cell.value)))
                          or (isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool))
                          for cell in others):
            return 'total'
    return None


def _mapping_issues(headers, mapping, asset_type, suggestion_ambiguous, explicit_mapping):
    blockers = []
    warnings = []
    duplicates = [header for header, count in Counter(headers).items() if count > 1]
    if duplicates:
        blockers.append({'code': 'duplicate_headers', 'headers': duplicates})
    schema = asset_type.schema or []
    known_fields = {field['key'] for field in schema}
    invalid = [header for header, key in mapping.items()
               if not isinstance(header, str) or (key is not None and (not isinstance(key, str) or key not in known_fields
                                                                      or header not in headers))]
    keys = [key for key in mapping.values() if key is not None and isinstance(key, str)]
    if invalid or len(keys) != len(set(keys)):
        blockers.append({'code': 'invalid_mapping', 'headers': [str(header) for header in invalid]})
    if not explicit_mapping and suggestion_ambiguous:
        blockers.append({'code': 'ambiguous_mapping', 'headers': suggestion_ambiguous})
    missing = [field['key'] for field in schema if field.get('required') and field['key'] not in keys]
    if missing:
        blockers.append({'code': 'missing_required_mapping', 'fields': missing})
    if not keys:
        blockers.append({'code': 'no_mapped_columns'})
    unmatched = [header for header in headers if mapping.get(header) is None]
    if unmatched:
        warnings.append({'code': 'unmatched_columns', 'headers': unmatched})
    return blockers, warnings


def _convert_row(cells, headers, mapping, asset_type):
    schema = asset_type.schema or []
    fields = {field['key']: field for field in schema}
    converted = {}
    reasons = defaultdict(list)
    errors = []
    for index, header in enumerate(headers):
        key = mapping.get(header)
        if key not in fields:
            continue
        cell = cells[index] if index < len(cells) else Cell()
        if cell.kind in ('f', 'e'):
            reason = 'formula' if cell.kind == 'f' else 'excel_error'
            reasons[reason].append(key)
            errors.append(f"'{header}': Formülü/hatalı hücreyi kontrol edip sabit değer yazın.")
            continue
        try:
            converted[key] = import_field_value(cell.value, fields[key].get('type', 'text'),
                                               stock_quantity=asset_type.tracking_mode == 'stock' and key == 'quantity')
        except (ValueError, TypeError, OverflowError) as exc:
            reasons['invalid_value'].append(key)
            errors.append(f"'{header}': {exc}")
    missing = [field['key'] for field in schema if field.get('required')
               and is_blank(converted.get(field['key'])) and not any(field['key'] in keys for keys in reasons.values())]
    if missing:
        reasons['missing_required'].extend(missing)
    if not reasons:
        try:
            converted = validate_custom_data(schema, converted, asset_type.tracking_mode)
        except ValidationError as exc:
            invalid = exc.detail.get('custom_data', {})
            reasons['invalid_value'].extend(invalid.keys() if isinstance(invalid, dict) else [])
            errors.append(str(exc.detail))
    if not any(not is_blank(value) for value in converted.values()) and not reasons:
        reasons['unmapped_values'] = []
        errors.append('Satırda eşleştirilmiş ekipman bilgisi bulunamadı.')
    return converted, reasons, errors


def analyze_inventory(file: WorkbookSource, asset_types: Iterable[AssetType], *,
                      asset_type: AssetType | None = None, mapping: dict[str, str | None] | None = None,
                      sheet_name: str | None = None, header_row: int | None = None,
                      data_start_row: int | None = None, skip_rows: Iterable[str] = (),
                      exclude_row_numbers: Iterable[int] = (), require_review: bool = False,
                      template_mode: bool = False) -> InventoryAnalysis:
    """Return a reusable import plan; preview and writes consume exactly this plan."""
    asset_types = list(asset_types)
    workbook, metadata_by_path = _bounded_workbook(file)
    try:
        # Relationship aliases otherwise cause the same XML grid to be allocated
        # repeatedly even though the archive budget counted it only once.
        _validate_worksheet_relations(workbook, metadata_by_path)
        template_type = None
        if '_mapping' in workbook.sheetnames:
            identifier = str(workbook['_mapping'].cell(row=2, column=1).value or '').strip()
            template_type = next((item for item in asset_types if str(item.pk) == identifier), None)
        detected_template = ({'asset_type_id': str(template_type.pk), 'asset_type_name': template_type.name}
                             if template_type else None)
        sheets = _read_sheets(workbook, metadata_by_path)
    finally:
        workbook.close()
    sheet, chosen_header, headers, blockers = _layout(sheets, asset_types, sheet_name, header_row, template_type)
    start = data_start_row or (4 if (template_mode or template_type) and chosen_header == 1 else chosen_header + 1)
    if start <= chosen_header:
        raise InventoryImportError('İlk veri satırı başlık satırından sonra olmalıdır.')
    if start > len(sheet.rows) + 1:
        raise InventoryImportError('İlk veri satırı seçilen sayfada bulunamadı.')
    exclusions = set(exclude_row_numbers)
    if any(number < start or number > len(sheet.rows) for number in exclusions):
        raise InventoryImportError('Atlanacak satırlar seçilen veri aralığında olmalıdır.')
    skipped = []
    skipped_count = 0
    rows = []
    custom_skips = {value.strip().casefold() for value in skip_rows if value.strip()}
    for number in range(start, len(sheet.rows) + 1):
        cells = sheet.rows[number - 1]
        reason = 'custom' if number in exclusions else _row_skip_reason(
            number, cells, headers, sheet, custom_skips, mapping or {}, asset_type,
        )
        if reason:
            skipped_count += 1
            if len(skipped) < 100:
                first = next((cell.value for cell in cells if not is_blank(cell.value)), '')
                skipped.append({'row_number': number, 'reason': reason, 'label': str(first)[:200]})
        else:
            rows.append((number, cells))
    type_columns = [index for index, header in enumerate(headers) if normalized(header) in TYPE_HEADERS]
    type_names = {normalized(cells[index].value) for _, cells in rows for index in type_columns
                  if index < len(cells) and not is_blank(cells[index].value)}
    suggestions = []
    for item in asset_types:
        suggested, _ = _suggest_mapping(headers, item)
        matched = sum(key is not None for key in suggested.values())
        if not matched:
            continue
        required = {field['key'] for field in item.schema or [] if field.get('required')}
        confidence = 'high' if required.issubset(set(suggested.values())) and matched >= min(2, len(item.schema or [])) else 'review'
        suggestions.append({'asset_type_id': str(item.pk), 'asset_type_name': item.name,
                            'confidence': confidence, 'mapping': suggested, 'matched_fields': matched})
    suggestions.sort(key=lambda item: (-item['matched_fields'], item['confidence'] != 'high', item['asset_type_name'], item['asset_type_id']))
    best = suggestions[0] if suggestions else None
    top = [item for item in suggestions if best and item['matched_fields'] == best['matched_fields']
           and item['confidence'] == best['confidence']]
    named_types = [item for item in asset_types if len(type_names) == 1 and normalized(item.name) in type_names]
    named_type = named_types[0] if len(named_types) == 1 else None
    named_suggestion = next((item for item in suggestions if named_type and item['asset_type_id'] == str(named_type.pk)), None)
    if template_type:
        suggested_id = str(template_type.pk)
    elif named_suggestion and named_suggestion['confidence'] == 'high':
        suggested_id = named_suggestion['asset_type_id']
    else:
        suggested_id = best['asset_type_id'] if len(top) == 1 and best['confidence'] == 'high' else None
    chosen_type = asset_type or template_type or next((item for item in asset_types if str(item.pk) == suggested_id), None)
    warnings = []
    if chosen_type is None:
        blockers.append({'code': 'asset_type_selection_required'})
        if not suggestions:
            warnings.append({'code': 'no_type_suggestion'})
    suggested_mapping, ambiguous = _suggest_mapping(headers, chosen_type) if chosen_type else ({header: None for header in headers}, [])
    resolved_mapping = mapping if mapping is not None else suggested_mapping
    if chosen_type:
        new_blockers, new_warnings = _mapping_issues(headers, resolved_mapping, chosen_type, ambiguous, mapping is not None)
        blockers.extend(new_blockers)
        warnings.extend(new_warnings)
    if len(type_names) > 1:
        blockers.append({'code': 'mixed_asset_types', 'headers': [headers[index] for index in type_columns]})
    converted = {}
    row_errors = {}
    row_reasons = {}
    unique = next((field['key'] for field in chosen_type.schema or [] if field.get('is_unique_key')), None) if chosen_type else None
    identifiers = defaultdict(list)
    if chosen_type:
        for number, cells in rows:
            values, reasons, errors = _convert_row(cells, headers, resolved_mapping, chosen_type)
            converted[number] = values
            row_reasons[number] = reasons
            if errors:
                row_errors[number] = errors
            identifier = values.get(unique) if unique else None
            if not is_blank(identifier):
                identifiers[json.dumps(identifier, sort_keys=True)].append(number)
        for numbers in identifiers.values():
            if len(numbers) > 1:
                for number in numbers:
                    row_reasons[number]['duplicate_identifier'].append(unique)
                    row_errors.setdefault(number, []).append('Aynı tanımlayıcı dosyada birden fazla satırda var; hangi satırın alınacağını seçin.')
    review_rows = []
    review_numbers = []
    for number, reasons in row_reasons.items():
        if reasons:
            review_numbers.append(number)
            for reason, fields in reasons.items():
                if len(review_rows) < 50:
                    review_rows.append({'row_number': number, 'reason': reason, 'fields': fields})
    if any('formula' in reasons for reasons in row_reasons.values()):
        warnings.append({'code': 'formula_review_required'})
    if any('duplicate_identifier' in reasons for reasons in row_reasons.values()):
        warnings.append({'code': 'duplicate_identifier', 'fields': [unique]})
    if not rows:
        blockers.append({'code': 'no_data_rows'})
    analysis = {
        'version': 1, 'sheet_name': sheet.name, 'sheet_names': [item.name for item in sheets],
        'header_row': chosen_header, 'data_start_row': start,
        'layout_confidence': 'review' if any(item['code'] in ('sheet_selection_required', 'header_selection_required') for item in blockers) else 'high',
        'suggested_asset_type_id': suggested_id, 'analysis_asset_type_id': str(chosen_type.pk) if chosen_type else None,
        'analysis_mapping': resolved_mapping, 'suggested_mapping': suggested_mapping, 'type_suggestions': suggestions[:20],
        'counts': {'data': len(rows), 'skipped': skipped_count, 'review': len(review_numbers),
                   'ready': len(rows) - len(review_numbers) if chosen_type else 0},
        'blockers': blockers, 'warnings': warnings, 'skipped_rows': skipped,
        'review_rows': review_rows, 'review_row_numbers': review_numbers,
    }
    preview_rows = [[str(cell.value) if cell.value is not None else '' for cell in cells[:len(headers)]]
                    for _, cells in rows[:50]]
    preview = {'headers': headers, 'preview_rows': preview_rows,
               'preview_row_numbers': [number for number, _ in rows[:50]], 'total_rows': len(rows),
               'detected_template': detected_template, 'analysis': analysis}
    return InventoryAnalysis(preview, rows, converted, row_errors, skipped_count)
